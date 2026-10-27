# -*- coding: utf-8 -*-
"""Tests 1.27 (lot C) : inspection et formation (grille APC notée, compte rendu signé, recommandations créées et suivies,
plans de formation, animations avec liste de présence et attestations), patrimoine et matériel (bâtiments repris de
l'inventaire, tables-bancs et besoin, manuels par élève, matériel, besoins, dotations et PV de réception), examens et
concours (CEP : épreuves par défaut, centres, jurys, candidats repris des CM2, numéros de table et salles, notes à l'écran
et par Excel, délibération avec repêchage et ex aequo, validation, statistiques, convocation, relevé, PV, liste des
admis ; concours : rangs, palmarès, diplôme de lauréat).
Lancer : python3 -m unittest tests.test_v127"""
import os
import subprocess
import sys
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
PREFIXE = "GESCOL" if os.path.exists(os.path.join(RACINE, "GESCOL_PUBLIC.md")) else "SCOLARIS"

SCRIPT = r'''
import html, io, json, os, re, sys, tempfile
from datetime import date, timedelta
P = os.environ["PREFIXE_SCRIPT"]
os.environ[P + "_EDITION"] = os.environ["EDITION_SCRIPT"]; os.environ[P + "_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config
from app import app, q, ex
from werkzeug.security import generate_password_hash
import openpyxl
app.config["TESTING"] = True
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.headers.get("Location"), r.data[-2500:]); return r
def txt(r): return html.unescape(r.data.decode())
def texte_pdf(r):
    import pdfplumber
    assert r.mimetype == "application/pdf", (r.request.path, r.mimetype)
    with pdfplumber.open(io.BytesIO(r.data)) as d:
        return "\n".join(p.extract_text() or "" for p in d.pages).replace("\n", " "), len(d.pages)
def xlsx(r):
    wb = openpyxl.load_workbook(io.BytesIO(ok(r).data)); ws = wb.active
    return [[c.value for c in l] for l in ws.iter_rows()]
def Q(*a, **k):
    with app.app_context():
        return q(*a, **k)
def EX(*a, **k):
    with app.app_context():
        return ex(*a, **k)
JUR = config.EDITION in ("cs", "ddemp")
S = config.ANNEE_ACTIVE
AUJ = date.today()
EX("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
c = app.test_client()
ok(c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}), 302)
if JUR:
    EX("INSERT OR REPLACE INTO jur_structure (id, niveau, departement, cs, nom) VALUES (1, ?, 'Plateau', 'Pobè', 'CS Pobè')", (config.EDITION,))
    cs = EX("INSERT INTO jur_cs (departement, nom) VALUES ('Plateau', 'Pobè')")
    zid = EX("INSERT INTO jur_zones (cs_id, nom, ordre) VALUES (?, 'Zone 1', 1)", (cs,))
    j1 = EX("INSERT INTO jur_ecoles (cs_id, zone_id, code, nom, type, secteur, actif) VALUES (?,?,?,?,?,?,1)", (cs, zid, "E1", "EPP Towé", "EP", "public"))
    ok(c.post(f"/juridiction/ecoles-modules/{j1}/ouvrir"), 302)
    eco = Q("SELECT id FROM ecoles WHERE nom='EPP Towé'", un=True)["id"]
else:
    eco = 1
NOM_E = "EPP Towé" if JUR else (config.ECOLE.get("nom") or "")
cm2 = EX("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES ('CM2 A', 'CM2', 0, ?)", (eco,))
ci = EX("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES ('CI A', 'CI', 0, ?)", (eco,))
for i, (n, p_, s_, dn, cl) in enumerate([("AGBO", "Rita", "F", "2014-03-02", cm2), ("SOSSOU", "Marc", "M", "2014-05-10", cm2), ("ADJOVI", "Paul", "M", "2013-01-20", cm2),
                                         ("BIO", "Awa", "F", "2014-07-01", cm2)] + [(f"KPO{k}", "X", "M", "2020-01-01", ci) for k in range(5)]):
    EX("INSERT INTO eleves (matricule, nom, prenoms, sexe, date_naissance, classe_id, statut, date_inscription, annee) VALUES (?,?,?,?,?,?,'inscrit','2026-09-15',?)",
       (f"M{i}", n, p_, s_, dn, cl, S))

# ===================== 1. Inspection et formation
data = {"enseignant": "HOUNKPE Jean", "sexeEns": "M", "cours": "CM2", "date": AUJ.isoformat(), "visiteur": "ADJAI Rémi", "qualiteVisiteur": "Conseiller pédagogique",
        "typeVisite": "Visite de classe", "discipline": "Mathématiques", "seance": "SA 2 — la proportionnalité",
        "pointsForts": "Bonne préparation", "recommandations": "Varier les supports.\n- Exploiter davantage les productions des élèves"}
data.update({f"c{i}": v for i, v in enumerate([4, 3, 3, 3, 2, 3, 3, 4, 4, 3], 1)})
ok(c.post(f"/reg/inspection/visites/0?annee={S}", data=data), 302)
t = txt(ok(c.get(f"/reg/inspection/visites?annee={S}")))
assert "16" in t and "Très bien" in t, "note 32/40 -> 16 /20"
vid = Q("SELECT id FROM reg_inspection_visites", un=True)["id"]
s, _n = texte_pdf(ok(c.get(f"/reg/inspection/visites/{vid}/cr.pdf")))
assert "COMPTE RENDU DE VISITE DE CLASSE" in s and "HOUNKPE Jean" in s and "Démarche APC" in s and "Très bien" in s and "Varier les supports" in s, s[:2000]
rows = xlsx(c.get(f"/reg/inspection/visites.xlsx?annee={S}"))
assert rows[0][:4] == ["Niveau", "Entité", "Saison", "Enseignant(e)"] and rows[0][-2:] == ["Note (/20)", "Appréciation"] and rows[1][-2] == 16, (rows[0], rows[1])
assert "Préparation et fiche pédagogique" in rows[0] and "Climat de classe et inclusion" in rows[0]
assert "Créer les recommandations" in t
ok(c.post(f"/inspection/visite/{vid}/recommandations"), 302)
ok(c.post(f"/inspection/visite/{vid}/recommandations"), 302)          # pas de doublon
r = Q("SELECT * FROM reg_inspection_recommandations ORDER BY id")
assert len(r) == 2 and r[0]["destinataire"] == "HOUNKPE Jean" and r[1]["recommandation"] == "Exploiter davantage les productions des élèves" and r[0]["statut"] == "Non commencée"
EX("UPDATE reg_inspection_recommandations SET echeance=? WHERE id=?", ((AUJ - timedelta(days=5)).isoformat(), r[0]["id"]))
EX("UPDATE reg_inspection_recommandations SET statut='Réalisée' WHERE id=?", (r[1]["id"],))
t = txt(ok(c.get(f"/inspection/suivi?annee={S}")))
assert "Varier les supports" in t and "50 % réalisées" in t and "moyenne 16" in t, t[t.find("Bilan"):][:2000]
ok(c.post(f"/reg/inspection/plans/0?annee={S}", data={"theme": "Pédagogie de l'intégration", "publicCible": "Enseignants du CM", "dateDebut": AUJ.isoformat(),
          "participantsPrevus": "25", "cout": "150000", "statut": "Prévue"}), 302)
ok(c.post(f"/reg/inspection/animations/0?annee={S}", data={"type": "Animation pédagogique d’UP", "date": AUJ.isoformat(), "theme": "La lecture au CI",
          "animateur": "ADJAI Rémi", "lieu": "EPP Towé", "presentsH": "6", "presentsF": "9",
          "participants": "HOUNKPE Jean — EPP Towé / enseignant\nSOGLO Marie — EPP Issaba / directrice\nDOSSOU Luc", "compteRendu": "Bonne participation."}), 302)
aid = Q("SELECT id FROM reg_inspection_animations", un=True)["id"]
assert "15" in txt(ok(c.get(f"/reg/inspection/animations?annee={S}")))
s, _n = texte_pdf(ok(c.get(f"/reg/inspection/animations/{aid}/cr.pdf")))
assert "Liste de présence" in s and "SOGLO Marie" in s and "Bonne participation" in s
s, n = texte_pdf(ok(c.get(f"/reg/inspection/animations/{aid}/attestations.pdf")))
assert n == 3 and s.count("ATTESTATION DE PARTICIPATION") == 3 and "DOSSOU Luc" in s and "La lecture au CI" in s, (n, s[:800])

# ===================== 2. Patrimoine et matériel
EX("INSERT INTO pat_locaux (ecole_id, type, designation, murs, etat, annee_service) VALUES (?, 'salle_classe', 'Salle du CM2', 'dur', 'mauvais', 1998)", (eco,))
EX("INSERT INTO pat_locaux (ecole_id, type, designation, murs, etat) VALUES (?, 'latrines', 'Latrines (4 cabines)', 'semi_dur', 'bon')", (eco,))
EX("INSERT INTO pat_mobilier (ecole_id, categorie, classe_id, quantite, etat) VALUES (?, 'tb2', ?, 1, 'bon')", (eco, cm2))
assert "2 local(aux)" in txt(ok(c.get("/patrimoine-materiel/reprendre")))
ok(c.post("/patrimoine-materiel/reprendre"), 302)
b = Q("SELECT * FROM reg_patrimoine_batiments ORDER BY designation")
assert [(x["designation"], x["nature"], x["etat"], x["materiaux"]) for x in b] == [("Latrines (4 cabines)", "Latrines", "Bon", "Semi-dur"),
                                                                             ("Salle du CM2", "Salle de classe", "Mauvais", "Définitif")], [dict(x) for x in b]
tb = {x["cours"]: dict(x) for x in Q("SELECT * FROM reg_patrimoine_tables")}
assert tb["CM2"]["eleves"] == 4 and tb["CM2"]["tablesBon"] == 1 and tb["CI"]["eleves"] == 5
t = txt(ok(c.get(f"/reg/patrimoine/tables?annee={S}")))
rows = xlsx(c.get(f"/reg/patrimoine/tables.xlsx?annee={S}"))
assert rows[0][:4] == ["Niveau", "Entité", "Année scolaire", "Cours"] and rows[0][-1] == "Besoin en tables-bancs", rows[0]
bes = {l[3]: l[-1] for l in rows[1:]}
assert bes["CM2"] == 1 and bes["CI"] == 3, bes
ok(c.post(f"/reg/patrimoine/manuels/0?annee={S}", data={"cours": "CM2", "matiere": "Mathématiques", "titre": "Champions en maths", "nombre": "6", "bonEtat": "2"}), 302)
rows = xlsx(c.get(f"/reg/patrimoine/manuels.xlsx?annee={S}"))
assert rows[1][-1] == 0.5 and rows[1][rows[0].index("Élèves du cours")] == 4, rows
ok(c.post(f"/reg/patrimoine/materiel/0?annee={S}", data={"designation": "Ballons de football", "categorie": "Sportif", "quantite": "5", "etat": "Bon", "valeur": "25000"}), 302)
ok(c.post(f"/reg/patrimoine/besoins/0?annee={S}", data={"designation": "Tables-bancs", "categorie": "Mobilier", "quantite": "4", "priorite": "Haute", "coutEstime": "140000"}), 302)
ok(c.post(f"/reg/patrimoine/dotations/0?annee={S}", data={"date": AUJ.isoformat(), "designation": "Kits scolaires", "quantite": "50", "unite": "kit", "source": "Partenaire / ONG",
          "beneficiaire": NOM_E, "pvNumero": "PV-07", "receptionnaire": "AKPO Jules", "etatReception": "Conforme"}), 302)
did = Q("SELECT id FROM reg_patrimoine_dotations", un=True)["id"]
s, _n = texte_pdf(ok(c.get(f"/reg/patrimoine/dotations/{did}/pv.pdf")))
assert "PROCÈS-VERBAL DE RÉCEPTION" in s and "Kits scolaires" in s and "PV-07" in s and "AKPO Jules" in s and "livreur" in s.lower(), s[:1500]

# ===================== 3. Examens (CEP)
SES = "CEP 2027"
ok(c.post("/examens/session", data={"session": SES, "action": "epreuves"}), 302)
ep = Q("SELECT * FROM reg_examens_epreuves WHERE session=? ORDER BY ordre", (SES,))
assert len(ep) == 7 and ep[0]["epreuve"] == "Communication écrite (rédaction)" and ep[2]["epreuve"] == "Mathématiques"
EX("UPDATE reg_examens_epreuves SET coefficient=2 WHERE id=?", (ep[2]["id"],))
ok(c.post("/examens/session", data={"session": SES, "action": "epreuves"}), 302)
assert Q("SELECT COUNT(*) n FROM reg_examens_epreuves WHERE session=?", (SES,), un=True)["n"] == 7
ok(c.post(f"/reg/examens/centres/0?annee={S}", data={"session": SES, "centre": "Centre Pobè 1", "code": "POB1", "ecoleSiege": NOM_E, "chefCentre": "KOKO Albert",
          "salles": "2", "capacite": "3"}), 302)
ok(c.post(f"/reg/examens/jurys/0?annee={S}", data={"session": SES, "centre": "Centre Pobè 1", "president": "ALIA Bruno", "vicePresident": "OGOU Lin",
          "secretaire": "DADA Eve", "membres": "M1\nM2", "date": AUJ.isoformat()}), 302)
ok(c.post("/examens/session", data={"session": SES, "action": "candidats"}), 302)
cands = Q("SELECT * FROM reg_examens_candidats WHERE session=? ORDER BY nom", (SES,))
assert [x["nom"] for x in cands] == ["ADJOVI", "AGBO", "BIO", "SOSSOU"] and cands[0]["entite"] == NOM_E
EX("UPDATE reg_examens_candidats SET centre='Centre Pobè 1' WHERE session=?", (SES,))
ok(c.post("/examens/session", data={"session": SES, "action": "numeros"}), 302)
num = {x["nom"]: (x["numeroTable"], x["salle"]) for x in Q("SELECT * FROM reg_examens_candidats WHERE session=?", (SES,))}
assert num == {"ADJOVI": ("POB1-0001", "Salle 1"), "AGBO": ("POB1-0002", "Salle 1"), "BIO": ("POB1-0003", "Salle 1"), "SOSSOU": ("POB1-0004", "Salle 2")}, num
t = txt(ok(c.get(f"/examens/session?session={SES}")))
assert "4 candidat(s)" in t and "0 sans numéro" in t
# notes à l'écran : AGBO 15 partout, ADJOVI 9 partout (math 10), BIO absente, SOSSOU via Excel
ids = {x["nom"]: x["id"] for x in Q("SELECT id, nom FROM reg_examens_candidats WHERE session=?", (SES,))}
form = {"session": SES}
for e in ep:
    form[f"n_{ids['AGBO']}_{e['id']}"] = "15"
    form[f"n_{ids['ADJOVI']}_{e['id']}"] = "10" if e["id"] == ep[2]["id"] else "9"
    form[f"n_{ids['BIO']}_{e['id']}"] = "ABS"
ok(c.post("/examens/notes", data=form), 302)
rows = xlsx(c.get(f"/examens/notes.xlsx?session={SES}"))
assert rows[0][:5] == ["N° de table", "Nom", "Prénoms", "Sexe", "École"] and rows[0][5] == "Communication écrite (rédaction)" and len(rows[0]) == 12, rows[0]
lig = next(l for l in rows if l[1] == "SOSSOU")
lig[5:] = [15] * 7
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Notes"; ws.append(rows[0]); ws.append(lig)
bio = io.BytesIO(); wb.save(bio); bio.seek(0)
ok(c.post("/examens/notes/import", data={"session": SES, "fichier": (bio, "notes.xlsx")}, content_type="multipart/form-data"), 302)
assert json.loads(Q("SELECT notes FROM reg_examens_candidats WHERE id=?", (ids["SOSSOU"],), un=True)["notes"])[str(ep[0]["id"])] == 15
# délibération : seuil 10, repêchage 9 ; AGBO et SOSSOU ex aequo 1ers, ADJOVI repêché (9,25), BIO absente
ok(c.post("/examens/session", data={"session": SES, "action": "param", "seuil": "10", "repechage": "9"}), 302)
t = txt(ok(c.get(f"/examens/deliberation?session={SES}")))
assert "Admis (repêché)" in t and "9,25" in t and "Absent" in t, t[t.find("<th>Rang"):][:3000]
ok(c.post("/examens/deliberation", data={"session": SES}), 302)
d = {x["nom"]: (x["rang"], x["decision"], x["moyenne"]) for x in Q("SELECT * FROM reg_examens_candidats WHERE session=?", (SES,))}
assert d["AGBO"][:2] == (1, "Admis") and d["SOSSOU"][:2] == (1, "Admis") and d["ADJOVI"][:2] == (3, "Admis (repêché)") and d["BIO"][1] == "Absent", d
assert d["ADJOVI"][2] == 9.25
t = txt(ok(c.get(f"/examens/deliberation?session={SES}")))
assert "Ensemble" in t and "100" in t
s, _n = texte_pdf(ok(c.get(f"/examens/deliberation.pdf?session={SES}&centre=Centre Pobè 1")))
assert "PROCÈS-VERBAL DE DÉLIBÉRATION" in s and "ALIA Bruno" in s and "DADA Eve" in s and "POB1-0002" in s and "taux de réussite : 100" in s, s[:2500]
s, _n = texte_pdf(ok(c.get(f"/examens/admis.pdf?session={SES}")))
assert "LISTE DES CANDIDATS ADMIS" in s and "3 admis sur 3 présent(s)" in s and "BIO" not in s.split("présent(s)")[1], s[:1500]
s, _n = texte_pdf(ok(c.get(f"/reg/examens/candidats/{ids['AGBO']}/convocation.pdf")))
assert "CONVOCATION DU CANDIDAT" in s and "POB1-0002" in s and "Salle 1" in s and "KOKO Albert" in s, s[:1500]
s, _n = texte_pdf(ok(c.get(f"/reg/examens/candidats/{ids['ADJOVI']}/releve.pdf")))
assert "RELEVÉ DE NOTES" in s and "9,25" in s and "repêché" in s and "ALIA Bruno" in s, s[:1500]
rows = xlsx(c.get(f"/reg/examens/candidats.xlsx?annee={S}"))
assert rows[0][-2:] == ["Moyenne", "Décision"] and "Admis" in [l[-1] for l in rows[1:]], rows[:3]

# ===================== 4. Concours scolaires
ok(c.post(f"/reg/examens/concours/0?annee={S}", data={"intitule": "Dictée de la zone 1", "type": "Dictée", "cours": "CM2", "date": AUJ.isoformat(),
          "lieu": "EPP Towé", "organisateur": "UP Pobè 1", "bareme": "20", "statut": "Réalisé"}), 302)
for nom, pts, rang in (("AGBO", "18", ""), ("SOSSOU", "16,5", ""), ("BIO", "16,5", ""), ("ADJOVI", "12", "")):
    ok(c.post(f"/reg/examens/participants/0?annee={S}", data={"concours": "Dictée de la zone 1", "nom": nom, "prenoms": "X", "sexe": "F", "ecole": NOM_E,
              "cours": "CM2", "points": pts, "rang": rang}), 302)
coid = Q("SELECT id FROM reg_examens_concours", un=True)["id"]
s, _n = texte_pdf(ok(c.get(f"/reg/examens/concours/{coid}/palmares.pdf")))
assert "PALMARÈS DU CONCOURS" in s and re.search(r"1 AGBO X F .*18", s) and re.search(r"2 (SOSSOU|BIO) X", s) and re.search(r"4 ADJOVI X", s), s[:1500]
pid = Q("SELECT id FROM reg_examens_participants WHERE nom='SOSSOU'", un=True)["id"]
s, _n = texte_pdf(ok(c.get(f"/reg/examens/participants/{pid}/prix.pdf")))
assert "DIPLÔME DE LAURÉAT" in s and "2e prix" in s and "Dictée de la zone 1" in s, s[:1200]
t = txt(ok(c.get(f"/reg/examens/participants?annee={S}")))
assert "Rang retenu" in t
# menus
t = txt(ok(c.get("/examens/session")))
assert "Examens et concours" in t and "Inspection et formation" in t and "Patrimoine et matériel" in t
print("OK-" + config.EDITION)
'''


def lancer(edition):
    env = dict(os.environ)
    for k in ("SCOLARIS_EDITION", "SCOLARIS_DATA", "GESCOL_EDITION", "GESCOL_DATA"):
        env.pop(k, None)
    env["EDITION_SCRIPT"] = edition
    env["PREFIXE_SCRIPT"] = PREFIXE
    return subprocess.run([sys.executable, "-c", SCRIPT], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)


class TestLotC(unittest.TestCase):
    def verifier(self, ed):
        r = lancer(ed)
        self.assertIn("OK-" + ed, r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_privee(self):
        self.verifier("privee")

    def test_publique(self):
        self.verifier("publique")

    def test_cs(self):
        self.verifier("cs")


class TestCalculsLotC(unittest.TestCase):
    def test_grille_tables_deliberation(self):
        import m_examens as E
        import m_inspection as I
        import m_patrimoine_pvs as P
        self.assertEqual(I.note_grille({"c1": 4, "c2": 3}), 17.5)
        self.assertIsNone(I.note_grille({}))
        self.assertEqual([I.appreciation(n) for n in (16, 14, 12, 10, 9.9)], ["Très bien", "Bien", "Assez bien", "Passable", "Insuffisant"])
        self.assertEqual(I.paragraphes("a\n\n- b\n• c"), ["a", "b", "c"])
        self.assertEqual(I.participants("NOM P — École A\nX - Y\nSeul"), [("NOM P", "École A"), ("X", "Y"), ("Seul", "")])
        self.assertEqual(P.besoin_tables(45, 10, 2), 13)
        self.assertEqual(P.besoin_tables(10, 10, 2), 0)
        self.assertEqual(P.ratio_manuels(15, 30), 0.5)
        ep = [{"id": 1, "coefficient": 1, "bareme": 20}, {"id": 2, "coefficient": 2, "bareme": 20}]
        cands = [{"id": 1, "notes": '{"1": 12, "2": 12}'}, {"id": 2, "notes": '{"1": 12, "2": 12}'}, {"id": 3, "notes": '{"1": 6, "2": 10}'},
                 {"id": 4, "notes": '{"1": "ABS", "2": "ABS"}'}, {"id": 5, "notes": '{"1": 20}', "statut": "Absent"}]
        r = {x["id"]: x for x in E.deliberer(cands, ep, 10, 8)}
        self.assertEqual((r[1]["rang"], r[2]["rang"], r[3]["rang"]), (1, 1, 3))
        self.assertEqual(r[3]["moyenne"], 8.67)
        self.assertEqual(r[3]["decision"], "Admis (repêché)")
        self.assertEqual((r[4]["decision"], r[5]["decision"]), ("Absent", "Absent"))
        self.assertEqual(E.numeroter([{"id": 1, "entite": "B", "nom": "A"}, {"id": 2, "entite": "A", "nom": "Z"}], {"code": "X", "capacite": 1}),
                         [(2, "X-0001", "Salle 1"), (1, "X-0002", "Salle 2")])
        self.assertEqual(E.rangs_concours([{"id": 1, "points": 10}, {"id": 2, "points": 12}, {"id": 3, "points": 10}, {"id": 4, "rang": 1, "points": 2}]),
                         {4: 1, 2: 1, 1: 2, 3: 2})
        self.assertEqual(E.lire_note("abs"), "ABS")
        self.assertEqual(E.lire_note("12,5"), 12.5)


if __name__ == "__main__":
    unittest.main()
