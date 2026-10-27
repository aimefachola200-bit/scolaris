# -*- coding: utf-8 -*-
"""Tests 1.26 (lot B) : communication et alertes (notes, circulaires, convocations diffusées aux écoles, boîte de
réception et accusés de réception, état des accusés signé, échéances et rappels, messages aux parents WhatsApp / SMS,
agenda, alertes) ; santé, cantine et inclusion (visites médicales -> aptitude des licences, IMC, certificat d'aptitude,
infirmerie, cantine : repas et stocks avec soldes et alerte, besoins spécifiques, scolarisation des filles, absences de
l'appel -> alerte de décrochage et fiche) ; Excel commun (Niveau · Entité · Saison, clé d'unicité à l'import).
Lancer : python3 -m unittest tests.test_v126"""
import os
import subprocess
import sys
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
PREFIXE = "GESCOL" if os.path.exists(os.path.join(RACINE, "GESCOL_PUBLIC.md")) else "SCOLARIS"

SCRIPT = r'''
import html, io, os, re, sys, tempfile
from datetime import date, timedelta
P = os.environ["PREFIXE_SCRIPT"]
os.environ[P + "_EDITION"] = os.environ["EDITION_SCRIPT"]; os.environ[P + "_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config
from app import app, q, ex
from werkzeug.security import generate_password_hash
import openpyxl
app.config["TESTING"] = True
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.headers.get("Location"), r.data[-2500:]); return r
def txt(r): return html.unescape(r.data.decode())
def texte_pdf(r):
    import pdfplumber
    assert r.mimetype == "application/pdf", (r.request.path, r.mimetype)
    with pdfplumber.open(io.BytesIO(r.data)) as d:
        return "\n".join(p.extract_text() or "" for p in d.pages).replace("\n", " ")
def xlsx(r):
    wb = openpyxl.load_workbook(io.BytesIO(ok(r).data)); ws = wb.active
    return [[c.value for c in l] for l in ws.iter_rows()]
def classeur(entetes, lignes, feuille):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = feuille; ws.append(entetes)
    for l in lignes: ws.append(l)
    b = io.BytesIO(); wb.save(b); b.seek(0); return b
def Q(*a, **k):
    with app.app_context():
        return q(*a, **k)
JUR = config.EDITION in ("cs", "ddemp")
S = config.ANNEE_ACTIVE
AUJ = date.today()
with app.app_context():
    ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
c = app.test_client()
ok(c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}), 302)

def preparer_ecole(eco):
    with app.app_context():
        cl = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES ('CM1 A', 'CM1', 0, ?)", (eco,))
        ids = []
        for i, (n, p_, s_, dn) in enumerate([("AGBO", "Rita", "F", "2015-03-02"), ("SOSSOU", "Marc", "M", "2016-05-10"), ("ADJOVI", "Paul", "M", "2013-01-20")]):
            ids.append(ex("INSERT INTO eleves (matricule, nom, prenoms, sexe, date_naissance, classe_id, statut, date_inscription, annee) VALUES (?,?,?,?,?,?,'inscrit','2026-09-15',?)",
                          (f"M{eco}{i}", n, p_, s_, dn, cl, S)))
        for k in range(4):        # AGBO Rita : 4 jours d'absence = 8 demi-journées
            d = (AUJ - timedelta(days=k + 1)).isoformat()
            ex("INSERT OR IGNORE INTO appels (classe_id, date) VALUES (?, ?)", (cl, d))
            ex("INSERT INTO presences (eleve_id, date, statut) VALUES (?, ?, 'absent')", (ids[0], d))
        ex("INSERT INTO presences (eleve_id, date, statut) VALUES (?, ?, 'absent')", (ids[1], (AUJ - timedelta(days=2)).isoformat()))
    return cl

# ===================== 1. Communication : note émise par la structure (CS) ou par l'école
if JUR:
    with app.app_context():
        ex("INSERT OR REPLACE INTO jur_structure (id, niveau, departement, cs, nom) VALUES (1, ?, 'Plateau', 'Pobè', 'CS Pobè')", (config.EDITION,))
        cs = ex("INSERT INTO jur_cs (departement, nom) VALUES ('Plateau', 'Pobè')")
        zid = ex("INSERT INTO jur_zones (cs_id, nom, ordre) VALUES (?, 'Zone 1', 1)", (cs,))
        j1 = ex("INSERT INTO jur_ecoles (cs_id, zone_id, code, nom, type, secteur, actif) VALUES (?,?,?,?,?,?,1)", (cs, zid, "E1", "EPP Towé", "EP", "public"))
        j2 = ex("INSERT INTO jur_ecoles (cs_id, zone_id, code, nom, type, secteur, actif) VALUES (?,?,?,?,?,?,1)", (cs, zid, "E2", "EPP Issaba", "EP", "public"))
t = txt(ok(c.get("/reg/communication/notes")))
assert "Notes, circulaires et convocations" in t and "Boîte de réception" in t and "Accusés de réception" in t
dem = (AUJ + timedelta(days=3)).isoformat()
ok(c.post(f"/reg/communication/notes/0?annee={S}", data={"niveau": "CS" if JUR else "ECOLE", "entite": "CS Pobè" if JUR else "Mon école", "type": "Convocation",
          "reference": "N°012/CS", "date": AUJ.isoformat(), "objet": "Réunion des directeurs", "destNiveau": "ECOLE",
          "destinataires": "EPP Towé" if JUR else "", "texte": "Vous êtes convoqués.\nPrésence obligatoire.", "lieuReunion": "Salle de la CS",
          "dateReunion": dem, "heureReunion": "09:00", "echeance": dem, "signataire": "Le Chef de la circonscription scolaire"}), 302)
nid = Q("SELECT id FROM reg_communication_notes WHERE objet='Réunion des directeurs'", un=True)["id"]
s = texte_pdf(ok(c.get(f"/reg/communication/notes/{nid}/note.pdf")))
assert "CONVOCATION" in s and "Réunion des directeurs" in s and "Salle de la CS" in s and "Présence obligatoire" in s and "code" in s.lower(), s[:1500]
assert "0 / 1" in txt(ok(c.get(f"/reg/communication/notes?annee={S}")))
rows = xlsx(c.get(f"/reg/communication/notes.xlsx?annee={S}"))
assert rows[0][:6] == ["Niveau", "Entité", "Saison", "Type", "Référence (n°)", "Date"], rows[0]
assert rows[0][-1] == "Accusés de réception" and rows[1][0] == ("CS" if JUR else "ECOLE")
if JUR:
    # l'école destinataire ouvre l'application depuis la CS : la note est dans sa boîte, elle accuse réception
    preparer_ecole_cs = None
    ok(c.post(f"/juridiction/ecoles-modules/{j1}/ouvrir"), 302)
    t = txt(ok(c.get("/liaison/boite")))
    assert "Réunion des directeurs" in t and "Non lue" in t and "EPP Towé" in t, t[t.find("Boîte"):][:2000]
    al = txt(ok(c.get("/liaison/alertes")))
    assert "Note non lue" in al and "Convocation" in al
    ok(c.post(f"/liaison/accuser/{nid}"), 302)
    assert "Reçue" in txt(ok(c.get("/liaison/boite")))
    ok(c.post("/juridiction/ecole/fermer"), 302)
    # une autre école n'est pas destinataire
    ok(c.post(f"/juridiction/ecoles-modules/{j2}/ouvrir"), 302)
    assert "Réunion des directeurs" not in txt(ok(c.get("/liaison/boite")))
    ok(c.post("/juridiction/ecole/fermer"), 302)
    assert "1 / 1" in txt(ok(c.get(f"/reg/communication/notes?annee={S}")))
    s = texte_pdf(ok(c.get(f"/reg/communication/notes/{nid}/accuses.pdf")))
    assert "ACCUSÉS DE RÉCEPTION" in s.upper() and "EPP Towé" in s and "1 accusé(s) sur 1" in s, s[:1200]
else:
    # école : une circulaire de la CS arrive par import Excel (autre appareil) -> boîte de réception, accusé, export
    ent = rows[0][:-1]
    lig = ["CS", "CS Pobè", S, "Circulaire", "N°3/CS", AUJ.strftime("%d/%m/%Y"), "Rentrée des classes", "ECOLE", "", "Texte de la circulaire",
           "", "", "", "", "", "Le CCS"]
    ok(c.post(f"/reg/communication/notes/import?annee={S}", data={"fichier": (classeur(ent, [lig, lig], "Notes, circulaires et convocations"), "n.xlsx")},
              content_type="multipart/form-data"), 302)
    assert Q("SELECT COUNT(*) n FROM reg_communication_notes WHERE objet='Rentrée des classes'", un=True)["n"] == 1   # clé d'unicité : 2e ligne = mise à jour
    t = txt(ok(c.get("/liaison/boite")))
    assert "Rentrée des classes" in t and "CS Pobè" in t and "Réunion des directeurs" not in t
    rid = Q("SELECT id FROM reg_communication_notes WHERE objet='Rentrée des classes'", un=True)["id"]
    ok(c.post(f"/liaison/accuser/{rid}"), 302)
    acc = xlsx(c.get(f"/reg/communication/lectures.xlsx?annee={S}"))
    assert acc[0][:7] == ["Niveau", "Entité", "Saison", "Note", "Date", "Lecteur", "Observations"], acc[0]
    assert acc[1][0] == "ECOLE" and acc[1][3] == "N°3/CS — Rentrée des classes", acc[1]

# ===================== 2. Échéances, agenda, messages aux parents, alertes
ok(c.post(f"/reg/communication/echeances/0?annee={S}", data={"objet": "Statistiques de rentrée", "type": "Statistiques",
          "dateLimite": (AUJ - timedelta(days=2)).isoformat(), "rappelJours": "7", "statut": "À faire"}), 302)
ok(c.post(f"/reg/communication/echeances/0?annee={S}", data={"objet": "Rapport trimestriel", "type": "Rapport",
          "dateLimite": (AUJ + timedelta(days=40)).isoformat(), "rappelJours": "7", "statut": "À faire"}), 302)
t = txt(ok(c.get(f"/reg/communication/echeances?annee={S}")))
assert "En retard (2 j)" in t and "À venir" in t
t = txt(ok(c.get("/liaison/alertes")))
assert "Statistiques de rentrée" in t and "Rapport trimestriel" not in t
assert "alerte(s) à traiter" in txt(ok(c.get(f"/reg/communication/agenda?annee={S}")))
ok(c.post(f"/reg/communication/agenda/0?annee={S}", data={"date": dem, "heureDebut": "08:00", "heureFin": "12:00", "objet": "Animation pédagogique",
          "type": "Animation pédagogique", "lieu": "EPP Towé"}), 302)
_s = texte_pdf(ok(c.get(f"/reg/communication/agenda.pdf?annee={S}"))); assert "Animation" in _s and "AGENDA" in _s.upper(), _s[:1500]
if not JUR:
    ok(c.post(f"/reg/communication/messages/0?annee={S}", data={"date": AUJ.isoformat(), "objet": "Réunion des parents", "texte": "Samedi à 9 h",
              "canal": "WhatsApp", "envoye": "Non", "destinataires": "Tous les parents d’élèves"}), 302)
    t = ok(c.get(f"/reg/communication/messages?annee={S}")).data.decode()
    assert "https://wa.me/?text=R%C3%A9union%20des%20parents%20%E2%80%94%20Samedi%20%C3%A0%209%20h" in t and "sms:?body=" in t and "data-copier" in t
    # liste certifiée signée (QR)
    s = texte_pdf(ok(c.get(f"/reg/communication/messages.pdf?annee={S}")))
    assert "MESSAGES AUX PARENTS" in s and "Réunion des parents" in s

# ===================== 3. Santé, cantine, inclusion (école)
if JUR:
    ok(c.post(f"/juridiction/ecoles-modules/{j1}/ouvrir"), 302)
    with app.app_context():
        eco = Q("SELECT id FROM ecoles WHERE nom='EPP Towé'", un=True)["id"]
else:
    eco = 1
preparer_ecole(eco)
# licence de l'élève, puis visite médicale -> aptitude reportée (sans accents ni casse)
ok(c.post(f"/reg/sport/licences/0?annee={S}", data={"nom": "AGBO", "prenoms": "Rita", "sexe": "F", "date_naissance": "2015-03-02",
          "ecole": "EPP Towé" if JUR else "", "aptitude": "Non examiné"}), 302)
t = txt(ok(c.get("/liaison/alertes")))
assert "Décrochage" in t and "AGBO Rita" in t and "8 demi-journées" in t and "SOSSOU" not in t, t[t.find("Alertes"):][:2500]
assert "« Non examiné »" in t
ok(c.post(f"/reg/sante/visites/0?annee={S}", data={"niveau": "ECOLE", "entite": "EPP Towé" if JUR else "Mon école", "nom": "Agbo", "prenoms": "RITA",
          "sexe": "F", "cours": "CM1", "dateVisite": AUJ.isoformat(), "examinateur": "Infirmier HOUNSA", "poids": "32", "taille": "140",
          "vue": "Normale", "aptitude": "Apte avec réserve", "reserve": "Asthme : effort modéré"}), 302)
assert Q("SELECT aptitude, medecin FROM reg_sport_licences WHERE nom='AGBO'", un=True)["aptitude"] == "Apte avec réserve"
t = txt(ok(c.get(f"/reg/sante/visites?annee={S}")))
assert "16,3" in t or "16.3" in t, "IMC"
vid = Q("SELECT id FROM reg_sante_visites", un=True)["id"]
s = texte_pdf(ok(c.get(f"/reg/sante/visites/{vid}/aptitude.pdf")))
assert "CERTIFICAT D'APTITUDE" in s.upper() and "Infirmier HOUNSA" in s and "apte avec réserve" in s and "Asthme" in s, s[:1500]
assert "FICHE DE VISITE MÉDICALE" in texte_pdf(ok(c.get(f"/reg/sante/visites/{vid}/fiche.pdf"))).upper()
rows = xlsx(c.get(f"/reg/sante/visites.xlsx?annee={S}"))
assert rows[0][:5] == ["Niveau", "Entité", "Saison", "Nom", "Prénoms"] and rows[0][-1] == "IMC" and rows[1][-1] == 16.3, rows[:2]
# import : même nom + prénoms + date -> mise à jour (aptitude Inapte reportée sur la licence)
ent = rows[0][:-1]
l2 = list(rows[1][:-1]); l2[ent.index("Aptitude sportive")] = "Inapte"
ok(c.post(f"/reg/sante/visites/import?annee={S}", data={"fichier": (classeur(ent, [l2], "Visites médicales"), "v.xlsx")}, content_type="multipart/form-data"), 302)
assert Q("SELECT COUNT(*) n FROM reg_sante_visites", un=True)["n"] == 1
assert Q("SELECT aptitude FROM reg_sport_licences WHERE nom='AGBO'", un=True)["aptitude"] == "Inapte"
# infirmerie
ok(c.post(f"/reg/sante/infirmerie/0?annee={S}", data={"date": AUJ.isoformat(), "heure": "10:15", "nom": "SOSSOU", "prenoms": "Marc", "nature": "Blessure",
          "soins": "Désinfection", "orientation": "Retour en classe", "parentInforme": "Oui"}), 302)
# cantine : repas (total servis), stocks (solde, alerte)
ok(c.post(f"/reg/sante/cantine/0?annee={S}", data={"date": AUJ.isoformat(), "repas": "Déjeuner", "prevusG": "50", "prevusF": "45", "servisG": "48", "servisF": "44",
          "menu": "Riz sauce arachide"}), 302)
assert "92" in txt(ok(c.get(f"/reg/sante/cantine?annee={S}")))
assert "JOURNAL DE LA CANTINE" in texte_pdf(ok(c.get(f"/reg/sante/cantine.pdf?annee={S}"))).upper()
for d in (("Riz", "Entrée", "100", "B1"), ("Riz", "Sortie", "30", "S1"), ("Huile", "Entrée", "10", "B2"), ("Huile", "Sortie", "10", "S2")):
    ok(c.post(f"/reg/sante/stocks/0?annee={S}", data={"date": AUJ.isoformat(), "denree": d[0], "mouvement": d[1], "quantite": d[2],
              "unite": "kg" if d[0] == "Riz" else "litre", "piece": d[3]}), 302)
t = txt(ok(c.get("/sante-scolaire/stocks")))
assert "70" in t and "Épuisé" in t, t[t.find("Soldes"):][:1500]
assert "Stock épuisé" in txt(ok(c.get("/liaison/alertes")))
# besoins spécifiques, filles
ok(c.post(f"/reg/sante/besoins/0?annee={S}", data={"nom": "ADJOVI", "prenoms": "Paul", "sexe": "M", "cours": "CM1", "type": "Handicap visuel",
          "amenagements": "Place au premier rang, documents agrandis", "statut": "En cours"}), 302)
bid = Q("SELECT id FROM reg_sante_besoins", un=True)["id"]
assert "Place au premier rang" in texte_pdf(ok(c.get(f"/reg/sante/besoins/{bid}/inclusion.pdf")))
t = txt(ok(c.get(f"/reg/sante/filles?annee={S}")))
assert 'name="inscrites" inputmode="numeric" value="1"' in t and 'name="garcons" inputmode="numeric" value="2"' in t, "effectifs proposés"
ok(c.post(f"/reg/sante/filles/0?annee={S}", data={"inscrites": "45", "garcons": "50", "abandons": "2", "club": "Oui"}), 302)
assert "0.9" in txt(ok(c.get(f"/reg/sante/filles?annee={S}")))
assert xlsx(c.get(f"/reg/sante/filles.xlsx?annee={S}"))[0][:3] == ["Niveau", "Entité", "Année scolaire"]
# absences 30 jours -> fiche de décrochage
t = txt(ok(c.get("/sante-scolaire/absences")))
assert "AGBO Rita" in t and "8" in t and "Créer la fiche" in t
eid = Q("SELECT e.id FROM eleves e WHERE e.nom='AGBO' AND e.classe_id IN (SELECT id FROM classes WHERE ecole_id=?)", (eco,), un=True)["id"]
ok(c.post("/sante-scolaire/absences", data={"eleve": eid}), 302)
x = Q("SELECT * FROM reg_sante_decrochage", un=True)
assert x["nom"] == "AGBO" and x["absences"] == 8 and x["statut"] == "Signalé"
s = texte_pdf(ok(c.get(f"/reg/sante/decrochage/{x['id']}/signalement.pdf")))
assert "SIGNALEMENT" in s.upper() and "AGBO" in s
assert "fiche ouverte" in txt(ok(c.get("/liaison/alertes")))
assert "seuil" in txt(ok(c.get("/sante-scolaire/absences?seuil=10"))).lower()
assert 'name="eleve"' not in txt(ok(c.get("/sante-scolaire/absences?seuil=10")))
# synthèse pour le pilotage
import m_pilotage  # noqa: F401
t = txt(ok(c.get(f"/pilotage?annee={S}")))
# réinitialiser, supprimer
ok(c.post(f"/reg/sante/infirmerie/reinitialiser?annee={S}"), 302)
assert Q("SELECT COUNT(*) n FROM reg_sante_infirmerie", un=True)["n"] == 0
ok(c.post(f"/reg/sante/besoins/{bid}/supprimer?annee={S}"), 302)
assert Q("SELECT COUNT(*) n FROM reg_sante_besoins", un=True)["n"] == 0
# menus
t = txt(ok(c.get("/reg/sante/visites")))
assert "Santé, cantine et inclusion" in t and "Communication et alertes" in t
print("OK-" + config.EDITION)
'''


def lancer(edition):
    env = dict(os.environ)
    for k in ("SCOLARIS_EDITION", "SCOLARIS_DATA", "GESCOL_EDITION", "GESCOL_DATA"):
        env.pop(k, None)
    env["EDITION_SCRIPT"] = edition
    env["PREFIXE_SCRIPT"] = PREFIXE
    return subprocess.run([sys.executable, "-c", SCRIPT], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)


class TestLotB(unittest.TestCase):
    def verifier(self, ed):
        r = lancer(ed)
        self.assertIn("OK-" + ed, r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_privee(self):
        self.verifier("privee")

    def test_publique(self):
        self.verifier("publique")

    def test_cs(self):
        self.verifier("cs")


class TestCalculsLotB(unittest.TestCase):
    def test_imc_soldes_echeances(self):
        from datetime import date
        import m_communication as MC
        import m_sante_scolaire as MS
        self.assertEqual(MS.imc(32, 140), 16.3)
        self.assertEqual(MS.imc("", 140), "")
        self.assertEqual(MS.soldes([{"denree": "Riz", "unite": "kg", "mouvement": "Entrée", "quantite": 10},
                                    {"denree": "Riz", "unite": "kg", "mouvement": "Sortie", "quantite": 12.5}]), {("Riz", "kg"): -2.5})
        j = date(2026, 10, 20)
        self.assertEqual(MC.etat_echeance({"dateLimite": "2026-10-18"}, j), "En retard (2 j)")
        self.assertEqual(MC.etat_echeance({"dateLimite": "2026-10-25", "rappelJours": 7}, j), "À rappeler (J-5)")
        self.assertEqual(MC.etat_echeance({"dateLimite": "2026-12-25"}, j), "À venir")
        self.assertEqual(MC.etat_echeance({"dateLimite": "2026-10-18", "statut": "Fait"}, j), "Fait")
        self.assertEqual(MC.libelle_note({"reference": "N°1", "objet": "Rentrée"}), "N°1 — Rentrée")
        import registres as R
        self.assertEqual(R.niveau_code("Circonscription scolaire"), "CS")
        self.assertEqual(R.niveau_code("École"), "ECOLE")
        self.assertEqual(R.niveau_code("zone"), "ZONE")


if __name__ == "__main__":
    unittest.main()
