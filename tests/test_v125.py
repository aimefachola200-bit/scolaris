# -*- coding: utf-8 -*-
"""Tests 1.25 (lot A) : colonnes Excel de la saison UASCEP identiques au module commun (export / import, ancien schéma
renommé), cachets vert et noir ; sport scolaire (licences et catégories, compétitions, tirage, calendrier, feuilles de match,
classements et départage, phase finale, performances et records, palmarès, qualifiés, diplômes signés, publication dans la
saison UASCEP) ; tableau de bord du pilotage (école, CS) ; moteur des registres (ajouter, modifier, supprimer,
réinitialiser, Excel, import, PDF signé).
Lancer : python3 -m unittest tests.test_v125"""
import os
import subprocess
import sys
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
PREFIXE = "GESCOL" if os.path.exists(os.path.join(RACINE, "GESCOL_PUBLIC.md")) else "SCOLARIS"

SCRIPT = r'''
import html, io, os, re, sys, tempfile
P = os.environ["PREFIXE_SCRIPT"]
os.environ[P + "_EDITION"] = os.environ["EDITION_SCRIPT"]; os.environ[P + "_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config
from app import app, q, ex
from werkzeug.security import generate_password_hash
app.config["TESTING"] = True
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.headers.get("Location"), r.data[-2500:]); return r
def txt(r): return html.unescape(r.data.decode())
def texte_pdf(r):
    import pdfplumber
    assert r.mimetype == "application/pdf", (r.request.path, r.mimetype)
    with pdfplumber.open(io.BytesIO(r.data)) as d:
        return "\n".join(p.extract_text() or "" for p in d.pages).replace("\n", " "), len(d.pages)
def code_de(s):
    return re.search(r"code ([0-9A-F]{10})", s).group(1)
import openpyxl
JUR = config.EDITION in ("cs", "ddemp")
S = config.ANNEE_ACTIVE
with app.app_context():
    ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
c = app.test_client()
ok(c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}), 302)
if JUR:
    with app.app_context():
        ex("INSERT OR REPLACE INTO jur_structure (id, niveau, departement, cs, nom) VALUES (1, ?, 'Plateau', 'Pobè', 'CS Pobè')", (config.EDITION,))
        cs = ex("INSERT INTO jur_cs (departement, nom) VALUES ('Plateau', 'Pobè')")
        zid = ex("INSERT INTO jur_zones (cs_id, nom, ordre) VALUES (?, 'Zone 1', 1)", (cs,))
        ex("INSERT INTO jur_up (cs_id, zone_id, nom, rup_nom) VALUES (?, ?, 'UP Pobè 1', 'ADJAI Rémi')", (cs, zid))
    niveau, entite = "UP", "UP Pobè 1"
else:
    niveau, entite = "ECOLE", "1"
    with app.app_context():
        cl = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES ('CM1 A', 'CM1', 0, 1)")
        for i, (n, p_, s_, dn) in enumerate([("AGBO", "Rita", "F", "2015-03-02"), ("SOSSOU", "Marc", "M", "2016-05-10"), ("ADJOVI", "Paul", "M", "2013-01-20")]):
            ex("INSERT INTO eleves (matricule, nom, prenoms, sexe, date_naissance, classe_id, statut, date_inscription, annee) VALUES (?,?,?,?,?,?,'inscrit','2026-09-15',?)",
               (f"M{i}", n, p_, s_, dn, cl, S))

# ---------------- 0. UASCEP : saison alignée sur le module commun
for i, (nom, fon) in enumerate([("AKPO", "Président"), ("BIO", "Secrétaire général"), ("DOSSOU", "Trésorier général"), ("OLA", "Responsable sport")], 1):
    ok(c.post(f"/uascep/membre/0?niveau={niveau}&entite={entite}", data={"ordre": i, "nom": nom, "prenoms": "X", "sexe": "M", "fonction": fon}), 302)
with app.app_context():
    b = q("SELECT * FROM uascep_bureaux WHERE niveau=? AND entite=? AND statut='actif'", (niveau, entite), un=True)
    pres = q("SELECT id FROM uascep_membres WHERE bureau_id=? AND fonction='Président'", (b["id"],), un=True)["id"]
bid = b["id"]
ok(c.post(f"/uascep/membre/{pres}/cachet", data={"cachet_mode": "genere", "cachet_forme": "rond", "cachet_couleur": "vert"}), 302)
assert "#166534" in txt(ok(c.get(f"/uascep/membre/{pres}/cachet")))
base = f"/uascep/bureau/{bid}/saison"
ok(c.post(f"{base}/caisse/0?saison={S}", data={"date": "2027-01-10", "nature": "Recette", "categorie": "Cotisations des écoles", "libelle": "Cotisation", "montant": "5000"}), 302)
ok(c.post(f"{base}/engagements/0?saison={S}", data={"ecole": "EPP Towé", "discipline": "Football", "jg": "12", "jf": "0", "encadreur": "SOSSA", "paye": "Oui"}), 302)
wb = openpyxl.load_workbook(io.BytesIO(ok(c.get(f"/uascep/bureau/{bid}/saison.xlsx?saison={S}")).data))
assert [x.value for x in wb["Activités"][1]][:15] == ["Date", "Discipline", "Catégorie", "Genre", "Phase", "Lieu / site", "Écoles", "Équipes", "Participants G",
                                                   "Participants F", "Total", "Encadreurs", "Résultats / classement", "Qualifiés au niveau supérieur", "Observations"]
assert [x.value for x in wb["Caisse"][1]][:9] == ["Date", "Nature", "Catégorie", "Libellé", "École (cotisation)", "Pièce justificative n°", "Recette (F)", "Dépense (F)", "Solde (F)"]
assert [x.value for x in wb["Engagements"][1]][:13] == ["École", "Discipline", "Catégorie", "Genre", "Nom de l’équipe", "Joueurs G", "Joueuses F", "Total", "Encadreur",
                                                      "Contact", "Droit d’engagement (F)", "Payé", "Date d’engagement"]
assert wb["Engagements"][2][7].value == 12 and wb["Caisse"][2][6].value == 5000
# fichier « JS » (mêmes feuilles) importé : caisse (recette / dépense sans montant), matériel
wb2 = openpyxl.Workbook(); ws = wb2.active; ws.title = "Caisse"
ws.append(["Date", "Nature", "Catégorie", "Libellé", "École (cotisation)", "Pièce justificative n°", "Recette (F)", "Dépense (F)", "Solde (F)"])
ws.append(["12/02/2027", "Dépense", "Transport", "Transport des équipes", "", "D9", "", 2000, 3000]); ws.append(["Total", "", "", "", "", "", 5000, 2000, ""])
ws2 = wb2.create_sheet("Matériel"); ws2.append(["Désignation", "Discipline", "Quantité", "État", "Origine"]); ws2.append(["Ballon", "Football", 4, "Bon", "Don"])
buf = io.BytesIO(); wb2.save(buf); buf.seek(0)
t = txt(ok(c.post(f"{base}/import?saison={S}", data={"fichier": (buf, "saison.xlsx")}, content_type="multipart/form-data", follow_redirects=True)))
assert "Caisse : 1" in t and "Matériel : 1" in t, t[-800:]
with app.app_context():
    d = q("SELECT * FROM uascep_caisse WHERE bureau_id=? AND piece='D9'", (bid,), un=True)
    assert d["nature"] == "Dépense" and d["montant"] == 2000
    assert q("SELECT origine FROM uascep_materiel WHERE bureau_id=?", (bid,), un=True)["origine"] == "Don"

# ---------------- 0 bis. logo UASCEP de l'entité (définition du sigle, texte du bas, couleurs, image importée)
from urllib.parse import quote
t = txt(ok(c.get(f"/uascep/logo/{niveau}/{quote(entite)}")))
assert "Logo de l'UASCEP" in t and "<svg" in t and "UNION DES" not in t[:0] and "Sportives" in t
ok(c.post(f"/uascep/logo/{niveau}/{quote(entite)}", data={"haut": "UASCEP", "definition": "Union des Associations Sportives et Culturelles des Écoles Primaires",
          "bas": "COMMUNE DE POBÈ", "c1": "#123456", "c2": "#E8B10E", "c3": "#1B8A3C"}), 302)
r = ok(c.get(f"/uascep/logo/{niveau}/{quote(entite)}.svg"))
assert r.mimetype == "image/svg+xml" and b"#123456" in r.data and "COMMUNE DE POBÈ".encode() in r.data and b'fill="#E8B10E" font-weight="900">S<' in r.data
r = ok(c.get(f"/uascep/logo/{niveau}/{quote(entite)}.png?taille=200"))
from PIL import Image
im = Image.open(io.BytesIO(r.data)); assert im.size == (200, 200)
s_, _ = texte_pdf(ok(c.get(f"/uascep/bureau/{bid}/composition.pdf")))
assert "ÉCOLES PRIMAIRES — UASCEP" in s_
ok(c.post(f"/uascep/logo/{niveau}/{quote(entite)}", data={"reinitialiser": "1"}), 302)
assert b"#123456" not in ok(c.get(f"/uascep/logo/{niveau}/{quote(entite)}.svg")).data

# ---------------- 1. sport scolaire : licences
t = txt(ok(c.get("/sport")))
assert "Sport scolaire" in t and "Licences" in t
if not JUR:
    ok(c.post("/sport/licences/eleves", data={"annee": S, "classes": "CM1 A"}), 302)
    with app.app_context():
        ls = q("SELECT * FROM reg_sport_licences ORDER BY id")
        assert len(ls) == 3 and all(l["numero"].startswith("LIC-") for l in ls)
        cats = {l["nom"]: l["categorie"] for l in ls}
        assert cats["SOSSOU"] == "Benjamins" and cats["AGBO"] == "Minimes" and cats["ADJOVI"] == "Cadets", cats
        assert ls[0]["numero"] == "LIC-" + S[2:4] + "-1-0001", ls[0]["numero"]
else:
    ok(c.post(f"/reg/sport/licences/0?annee={S}", data={"nom": "KOUTON", "prenoms": "Ali", "sexe": "M", "date_naissance": "2014-02-01", "ecole": "EPP Towé"}), 302)
    with app.app_context():
        l = q("SELECT * FROM reg_sport_licences", un=True)
        assert l["categorie"] == "Minimes" and l["numero"].startswith("LIC-")
t = txt(ok(c.get(f"/reg/sport/licences?annee={S}")))
assert "Licences des élèves" in t and "Âge (saison)" in t
s_, _ = texte_pdf(ok(c.get(f"/reg/sport/licences.pdf?annee={S}")))
assert "REGISTRE DES LICENCES" in s_ and "Document vérifiable" in s_
assert "registre des licences" in txt(ok(c.get("/verifier/" + code_de(s_)))).lower()
with app.app_context():
    lid = q("SELECT id FROM reg_sport_licences", un=True)["id"]
s_, _ = texte_pdf(ok(c.get(f"/reg/sport/licences/{lid}.pdf")))
assert "CARTE DE LICENCE" in s_
wb = openpyxl.load_workbook(io.BytesIO(ok(c.get(f"/reg/sport/licences.xlsx?annee={S}")).data))
assert wb.active.title == "Licences" and wb.active["E1"].value == "Nom" and wb.active["A2"].value == "ECOLE"
# moteur : modifier, import, réinitialiser
ok(c.post(f"/reg/sport/licences/{lid}?annee={S}", data={"nom": "MODIFIE", "prenoms": "Z", "sexe": "F", "categorie": "Cadets"}), 302)
with app.app_context():
    assert q("SELECT nom, categorie FROM reg_sport_licences WHERE id=?", (lid,), un=True)["categorie"] == "Cadets"
wb3 = openpyxl.load_workbook(io.BytesIO(ok(c.get("/reg/sport/officiels/modele.xlsx")).data)); ws = wb3.active
ws.append(["HOUNSOU Jean", "Arbitre", "Football", "EPP A", "0197", "", ""])
buf = io.BytesIO(); wb3.save(buf); buf.seek(0)
ok(c.post(f"/reg/sport/officiels/import?annee={S}", data={"fichier": (buf, "o.xlsx")}, content_type="multipart/form-data"), 302)
with app.app_context():
    assert q("SELECT qualite FROM reg_sport_officiels", un=True)["qualite"] == "Arbitre"
ok(c.post(f"/reg/sport/officiels/reinitialiser?annee={S}"), 302)
with app.app_context():
    assert not q("SELECT 1 FROM reg_sport_officiels", un=True)

# ---------------- compétition par poules puis phase finale
ok(c.post(f"/sport/competitions?saison={S}", data={"niveau": niveau, "entite": entite, "intitule": "Coupe de football", "discipline": "Football",
          "categorie": "Minimes", "genre": "Garçons", "formule": "poules", "nb_poules": "2", "qualifies_poule": "2", "qualifies_sup": "2",
          "date_debut": "2027-03-01", "lieu": "Stade"}), 302)
with app.app_context():
    cid = q("SELECT id FROM sport_competitions WHERE intitule='Coupe de football'", un=True)["id"]
for n in ("EPP A", "EPP B", "EPP C", "EPP D", "EPP E", "EPP F", "EPP G", "EPP H"):
    ok(c.post(f"/sport/competition/{cid}/equipe/0", data={"nom": n, "ecole": n, "effectif_g": "12", "tete_serie": "1" if n in ("EPP A", "EPP B") else ""}), 302)
ok(c.post(f"/sport/competition/{cid}/equipes/importer", data={"source": "engagements"}), 302)   # EPP Towé (fiche d'engagement)
with app.app_context():
    assert q("SELECT COUNT(*) n FROM sport_equipes WHERE competition_id=?", (cid,), un=True)["n"] == 9
    ex("DELETE FROM sport_equipes WHERE competition_id=? AND nom='EPP Towé'", (cid,))
ok(c.post(f"/sport/competition/{cid}/tirage"), 302)
with app.app_context():
    eqs = q("SELECT * FROM sport_equipes WHERE competition_id=?", (cid,))
    assert {e["poule"] for e in eqs} == {"A", "B"} and sum(1 for e in eqs if e["poule"] == "A") == 4
    assert {e["poule"] for e in eqs if e["tete_serie"]} == {"A", "B"}
    ms = q("SELECT * FROM sport_matchs WHERE competition_id=?", (cid,))
    assert len(ms) == 12, len(ms)
    nom = {e["id"]: e["nom"] for e in eqs}
# résultats : l'équipe au plus petit nom (ordre alphabétique) gagne
import itertools
for m in ms:
    a, b_ = nom[m["equipe_a"]], nom[m["equipe_b"]]
    ok(c.post(f"/sport/competition/{cid}/match/{m['id']}", data={"score_a": "2" if a < b_ else "0", "score_b": "0" if a < b_ else "2", "arbitre": "HOUNSOU"}), 302)
t = txt(ok(c.get(f"/sport/competition/{cid}?onglet=classement")))
assert "Poule A" in t and "Poule B" in t
ok(c.post(f"/sport/competition/{cid}/phase-finale"), 302)
with app.app_context():
    demi = q("SELECT * FROM sport_matchs WHERE competition_id=? AND phase='elimination'", (cid,))
    assert len(demi) == 2 and demi[0]["tour"] == "Demi-finales"
for m in demi:
    ok(c.post(f"/sport/competition/{cid}/match/{m['id']}", data={"score_a": "1", "score_b": "1", "tab_a": "4", "tab_b": "3"}), 302)
ok(c.post(f"/sport/competition/{cid}/phase-finale"), 302)
with app.app_context():
    fin = q("SELECT * FROM sport_matchs WHERE competition_id=? AND tour IN ('Finale', 'Match pour la 3e place')", (cid,))
    assert len(fin) == 2
for m in fin:
    ok(c.post(f"/sport/competition/{cid}/match/{m['id']}", data={"score_a": "3", "score_b": "1"}), 302)
t = txt(ok(c.get(f"/sport/competition/{cid}?onglet=resultats")))
assert "Podium" in t
ok(c.post(f"/sport/competition/{cid}/publier", data={"terminer": "1"}), 302)
with app.app_context():
    a_ = q("SELECT * FROM uascep_activites WHERE source=?", (f"sport:{cid}",), un=True)
    assert a_ and a_["equipes"] == 8 and a_["pg"] == 96 and "1e" in a_["resultats"] and a_["statut"] == "Réalisée", dict(a_)
s_, _ = texte_pdf(ok(c.get(f"/uascep/bureau/{bid}/rapport.pdf?saison={S}")))
assert "Coupe de football" in s_
for doc, attendu in (("calendrier", "CALENDRIER DES RENCONTRES"), ("classement", "CLASSEMENT"), ("resultats", "RÉSULTATS OFFICIELS")):
    s_, _ = texte_pdf(ok(c.get(f"/sport/competition/{cid}/{doc}.pdf")))
    assert attendu in s_ and "Document vérifiable" in s_, (doc, s_[:300])
    assert "Sport scolaire" in txt(ok(c.get("/verifier/" + code_de(s_))))
s_, _ = texte_pdf(ok(c.get(f"/sport/competition/{cid}/match/{fin[0]['id']}.pdf")))
assert "FEUILLE DE MATCH" in s_ and "arbitre" in s_.lower()
s_, n = texte_pdf(ok(c.get(f"/sport/competition/{cid}/diplomes.pdf?type=laureats")))
assert "DIPLÔME D'HONNEUR" in s_ and "CHAMPION" in s_ and n >= 3 and "PRIMAIRES (UASCEP)" in s_, (n, s_[:300])
s_, n = texte_pdf(ok(c.get(f"/sport/competition/{cid}/diplomes.pdf?type=participation")))
assert n == 8 and "DE PARTICIPATION" in s_
t = txt(ok(c.get(f"/sport/palmares?saison={S}")))
assert "Coupe de football" in t
s_, _ = texte_pdf(ok(c.get(f"/sport/palmares.pdf?saison={S}")))
assert "PALMARÈS" in s_

# ---------------- athlétisme : performances, classement, records
ok(c.post(f"/sport/competitions?saison={S}", data={"niveau": niveau, "entite": entite, "intitule": "Cross et sprint", "discipline": "Athlétisme",
          "formule": "performances", "qualifies_sup": "2"}), 302)
with app.app_context():
    aid = q("SELECT id FROM sport_competitions WHERE intitule='Cross et sprint'", un=True)["id"]
for n, perf in (("A1", "9''8"), ("A2", "10,2"), ("A3", "9.5"), ("A4", "")):
    ok(c.post(f"/sport/competition/{aid}/perf/0", data={"epreuve": "60 m", "nom": n, "sexe": "M", "categorie": "Minimes", "performance": perf}), 302)
ok(c.post(f"/sport/competition/{aid}/perf/0", data={"epreuve": "Saut en longueur", "nom": "S1", "sexe": "F", "performance": "3 m 45"}), 302)
ok(c.post(f"/sport/competition/{aid}/perf/0", data={"epreuve": "Saut en longueur", "nom": "S2", "sexe": "F", "performance": "3,50 m"}), 302)
t = txt(ok(c.get(f"/sport/competition/{aid}?onglet=perfs")))
assert "record" in t
t = txt(ok(c.get("/sport/records")))
assert "9.5" in t and "3,50 m" in t
s_, n = texte_pdf(ok(c.get(f"/sport/competition/{aid}/diplomes.pdf?type=laureats")))
assert "A3" in s_ and n == 5, n
# qualification vers le niveau supérieur (CS seulement : plusieurs niveaux dans la même application)
if JUR:
    ok(c.post(f"/sport/competitions?saison={S}", data={"niveau": "ZONE", "entite": "Zone 1", "intitule": "Finale zone", "discipline": "Athlétisme", "formule": "performances"}), 302)
    with app.app_context():
        zc = q("SELECT id FROM sport_competitions WHERE intitule='Finale zone'", un=True)["id"]
    ok(c.post(f"/sport/competition/{aid}/qualifier", data={"cible": zc}), 302)
    with app.app_context():
        assert q("SELECT COUNT(*) n FROM sport_perfs WHERE competition_id=?", (zc,), un=True)["n"] == 4   # 2 au 60 m + 2 au saut

# ---------------- championnat aller-retour et forfait
ok(c.post(f"/sport/competitions?saison={S}", data={"niveau": niveau, "entite": entite, "intitule": "Handball", "discipline": "Handball", "formule": "championnat_ar"}), 302)
with app.app_context():
    hid = q("SELECT id FROM sport_competitions WHERE intitule='Handball'", un=True)["id"]
for n in ("H1", "H2", "H3"):
    ok(c.post(f"/sport/competition/{hid}/equipe/0", data={"nom": n}), 302)
ok(c.post(f"/sport/competition/{hid}/tirage"), 302)
with app.app_context():
    ms = q("SELECT * FROM sport_matchs WHERE competition_id=?", (hid,))
    assert len(ms) == 6
ok(c.post(f"/sport/competition/{hid}/match/{ms[0]['id']}", data={"forfait": "A"}), 302)
t = txt(ok(c.get(f"/sport/competition/{hid}?onglet=classement")))
assert ">3<" in t
r = c.post(f"/sport/competition/{hid}/tirage", follow_redirects=True)
assert "Des résultats sont déjà saisis" in txt(r)
ok(c.post(f"/sport/competition/{hid}/calendrier/reinitialiser"), 302)
with app.app_context():
    assert not q("SELECT 1 FROM sport_matchs WHERE competition_id=?", (hid,), un=True)
ok(c.post(f"/sport/competition/{hid}/supprimer"), 302)
ok(c.get(f"/sport/competition/{hid}"), 404)
ok(c.post("/sport/parametres", data={"cat_Poussins": "0-7", "cat_Benjamins": "8-10", "cat_Minimes": "11-12", "cat_Cadets": "13-99", "recalculer": "1"}), 302)
assert "0-7" in txt(ok(c.get("/sport/parametres")))
# classeur commun du sport : export puis réimport (mise à jour, pas de doublon)
wb = openpyxl.load_workbook(io.BytesIO(ok(c.get(f"/sport/saison.xlsx?saison={S}")).data))
for f_ in ("Licences", "Compétitions", "Équipes engagées", "Calendrier et feuilles de match", "Athlétisme (performances)", "Palmarès et qualifiés"):
    assert f_ in wb.sheetnames, (f_, wb.sheetnames)
assert [x.value for x in wb["Compétitions"][1]][:4] == ["Niveau", "Entité", "Saison", "Intitulé"]
assert [x.value for x in wb["Licences"][1]][:5] == ["Niveau", "Entité", "Saison", "N° de licence", "Nom"]
assert wb["Compétitions"][2][7].value == "Poules puis élimination directe"
with app.app_context():
    avant = [q(f"SELECT COUNT(*) n FROM {t}", un=True)["n"] for t in ("sport_competitions", "sport_equipes", "sport_matchs", "sport_perfs")]
buf = io.BytesIO(ok(c.get(f"/sport/saison.xlsx?saison={S}")).data)
t = txt(ok(c.post(f"/sport/import?saison={S}", data={"fichier": (buf, "s.xlsx")}, content_type="multipart/form-data", follow_redirects=True)))
assert "Import : Compétitions" in t, t[-1500:]
with app.app_context():
    apres = [q(f"SELECT COUNT(*) n FROM {t}", un=True)["n"] for t in ("sport_competitions", "sport_equipes", "sport_matchs", "sport_perfs")]
    assert avant == apres, (avant, apres)

# ---------------- 2. tableau de bord du pilotage
t = txt(ok(c.get(f"/pilotage?annee={S}")))
assert "Tableau de bord du pilotage" in t and "Ratio élèves / maître" in t and "Licences sportives" in t
if not JUR:
    assert "Effectifs par classe" in t and "CM1 A" in t
    with app.app_context():
        assert q("SELECT 1 FROM pilotage_instantanes WHERE annee=?", (S,), un=True)
s_, _ = texte_pdf(ok(c.get(f"/pilotage.pdf?annee={S}")))
assert "TABLEAU DE BORD DU PILOTAGE" in s_ and "Document vérifiable" in s_
wb = openpyxl.load_workbook(io.BytesIO(ok(c.get(f"/pilotage.xlsx?annee={S}")).data))
assert "Indicateurs" in wb.sheetnames
# saisie des indicateurs : la ligne saisie prime sur les données calculées
ecole_nom = "EPP Towé" if JUR else config.ECOLE.get("nom")
if JUR:
    with app.app_context():
        ex("INSERT INTO jur_ecoles (cs_id, zone_id, code, nom, type, secteur, actif) VALUES (?,?,?,?,?,?,1)", (cs, zid, "E1", "EPP Towé", "EP", "public"))
ok(c.post(f"/reg/pilotage/indicateurs/0?annee={S}", data={"ecole": ecole_nom, "eff_g": "120", "eff_f": "100", "enseignants": "4", "reussite_t1": "62,5",
          "cep_presents_g": "10", "cep_presents_f": "10", "cep_admis_g": "8", "cep_admis_f": "9"}), 302)
t = txt(ok(c.get(f"/pilotage?annee={S}")))
assert "220" in t and "62,5 %" in t and "85 %" in t, t[t.find("Indicateur"):][:3000]
wb = openpyxl.load_workbook(io.BytesIO(ok(c.get(f"/reg/pilotage/indicateurs.xlsx?annee={S}")).data))
assert [x.value for x in wb.active[1]][:5] == ["Niveau", "Entité", "Année scolaire", "École", "Effectif garçons"]
if not JUR:
    with app.app_context():
        import m_transmission  # noqa: F401
        from flask import g as _g
    with app.test_request_context():
        pass
print("OK-" + config.EDITION)
'''


def lancer(edition):
    env = dict(os.environ)
    for k in ("SCOLARIS_EDITION", "SCOLARIS_DATA", "GESCOL_EDITION", "GESCOL_DATA"):
        env.pop(k, None)
    env["EDITION_SCRIPT"] = edition
    env["PREFIXE_SCRIPT"] = PREFIXE
    return subprocess.run([sys.executable, "-c", SCRIPT], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)


class TestLotA(unittest.TestCase):
    def verifier(self, ed):
        r = lancer(ed)
        self.assertIn("OK-" + ed, r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_privee(self):
        self.verifier("privee")

    def test_publique(self):
        self.verifier("publique")

    def test_cs(self):
        self.verifier("cs")


class TestCalculsSport(unittest.TestCase):
    def test_perfs(self):
        import m_sport as S
        self.assertEqual(S.valeur_perf("9''8"), 9.8)
        self.assertEqual(S.valeur_perf("1'12''5"), 72.5)
        self.assertEqual(S.valeur_perf("3 m 45", "m"), 3.45)
        self.assertEqual(S.valeur_perf("3,5 m", "m"), 3.5)
        self.assertEqual(S.valeur_perf("14 pts", "pts"), 14.0)
        self.assertIsNone(S.valeur_perf(""))

    def test_categories(self):
        import m_sport as S
        self.assertEqual(S.categorie_age("2016", "2026-2027"), "Benjamins")
        self.assertEqual(S.categorie_age("2018", "2026-2027"), "Poussins")
        self.assertEqual(S.categorie_age("2015", "2026-2027"), "Minimes")
        self.assertEqual(S.categorie_age("2013", "2026-2027"), "Cadets")

    def test_robin(self):
        import m_sport as S
        m = S.tour_robin([1, 2, 3, 4])
        self.assertEqual(len(m), 6)
        self.assertEqual({frozenset((a, b)) for _j, a, b in m}, {frozenset(x) for x in ((1, 2), (1, 3), (1, 4), (2, 3), (2, 4), (3, 4))})
        self.assertEqual(len(S.tour_robin([1, 2, 3], True)), 6)

    def test_classement_departage(self):
        import m_sport as S
        eq = [{"id": i, "nom": n} for i, n in ((1, "A"), (2, "B"), (3, "C"))]
        ms = [{"equipe_a": 1, "equipe_b": 2, "score_a": 1, "score_b": 0}, {"equipe_a": 2, "equipe_b": 3, "score_a": 3, "score_b": 0},
              {"equipe_a": 3, "equipe_b": 1, "score_a": 1, "score_b": 0}]
        cl = S.classement(eq, ms)
        self.assertEqual([x["nom"] for x in cl], ["B", "A", "C"])   # 3 pts chacun : différence générale
        ms2 = [{"equipe_a": 1, "equipe_b": 2, "score_a": 1, "score_b": 0}, {"equipe_a": 1, "equipe_b": 3, "score_a": 0, "score_b": 5},
               {"equipe_a": 2, "equipe_b": 3, "score_a": 1, "score_b": 0}]
        cl = S.classement(eq, ms2)
        self.assertEqual(cl[0]["nom"], "C")
        cl = S.classement(eq[:2], [{"equipe_a": 1, "equipe_b": 2, "forfait": "A"}])
        self.assertEqual((cl[0]["nom"], cl[0]["pts"], cl[1]["pts"], cl[0]["bp"]), ("B", 3, 0, 3))

    def test_elimination(self):
        import m_sport as S
        p = S.tableau_elimination([1, 2, 3, 4, 5, 6])
        self.assertEqual(p[:2], [(1, None), (2, None)])
        self.assertEqual(len(p), 4)
        self.assertEqual(S.vainqueur({"equipe_a": 1, "equipe_b": 2, "score_a": 1, "score_b": 1, "tab_a": 5, "tab_b": 4}), 1)
        self.assertIsNone(S.vainqueur({"equipe_a": 1, "equipe_b": 2, "score_a": 1, "score_b": 1}))

    def test_poules(self):
        import random
        import m_sport as S
        eq = [{"id": i, "tete_serie": i <= 2} for i in range(1, 10)]
        rep = S.repartir_poules(eq, 2, random.Random(1))
        self.assertNotEqual(rep[1], rep[2])
        self.assertEqual(sorted(list(rep.values()).count(p) for p in "AB"), [4, 5])


# points de code Extended_Pictographic (Unicode 15) + sélecteur de variante : aucun émoji dans l'application
_EP = """00A9 00AE 203C 2049 2122 2139 2194-2199 21A9-21AA 231A-231B 2328 2388 23CF 23E9-23F3 23F8-23FA 24C2 25AA-25AB 25B6 25C0
25FB-25FE 2600-2605 2607-2612 2614-2685 2690-2705 2708-2712 2714 2716 271D 2721 2728 2733-2734 2744 2747 274C 274E 2753-2755
2757 2763-2767 2795-2797 27A1 27B0 27BF 2934-2935 2B05-2B07 2B1B-2B1C 2B50 2B55 3030 303D 3297 3299 1F000-1F0FF 1F10D-1F10F
1F12F 1F16C-1F171 1F17E-1F17F 1F18E 1F191-1F19A 1F1AD-1F1E5 1F201-1F20F 1F21A 1F22F 1F232-1F23A 1F23C-1F23F 1F249-1F3FA
1F400-1F53D 1F546-1F64F 1F680-1F6FF 1F774-1F77F 1F7D5-1F7FF 1F80C-1F80F 1F848-1F84F 1F85A-1F85F 1F888-1F88F 1F8AE-1F8FF
1F90C-1F93A 1F93C-1F945 1F947-1FAFF 1FC00-1FFFD FE0F 1F1E6-1F1FF"""
_PLAGES = [(int(a, 16), int(b or a, 16)) for a, _s, b in (x.partition("-") for x in _EP.split())]


def _emoji(ch):
    o = ord(ch)
    return any(a <= o <= b for a, b in _PLAGES)


class TestSansEmoji(unittest.TestCase):
    """Règle 1.25 : pas d'émojis, des icônes au trait (icones.py, {{ icone('nom') }})."""

    def test_aucun_emoji(self):
        fautes = []
        cibles = [os.path.join(RACINE, "templates"), os.path.join(RACINE, "static")]
        fichiers = [os.path.join(RACINE, f) for f in os.listdir(RACINE) if f.endswith(".py")]
        for d in cibles:
            for racine, _ds, fs in os.walk(d):
                fichiers += [os.path.join(racine, f) for f in fs if f.endswith((".html", ".js", ".css", ".py", ".json", ".webmanifest", ".svg"))]
        for p in fichiers:
            try:
                with open(p, encoding="utf-8") as f:
                    t = f.read()
            except (UnicodeDecodeError, OSError):
                continue
            for i, l in enumerate(t.splitlines(), 1):
                if any(_emoji(c) for c in l):
                    fautes.append(f"{os.path.relpath(p, RACINE)}:{i}")
        self.assertEqual(fautes, [], "émojis trouvés : " + ", ".join(fautes[:30]))

    def test_icones(self):
        import icones
        self.assertIn("<svg", icones.svg("telecharger"))
        self.assertIn('stroke="currentColor"', icones.svg("ecole", 20, "École"))
        self.assertIn('aria-label="École"', icones.svg("ecole", 20, "École"))
        self.assertGreater(len(icones.D), 80)
        with open(os.path.join(RACINE, "templates", "base.html"), encoding="utf-8") as f:
            self.assertIn("icone(ico)", f.read())


class TestResumesProgression(unittest.TestCase):
    """1.25 : résumés de Français (outils de la langue) et de Mathématiques (domaines) rattachés à la progression annuelle."""

    @classmethod
    def setUpClass(cls):
        import gzip
        import json
        with gzip.open(os.path.join(RACINE, "static", "pedafiche_corpus.json.gz"), "rt", encoding="utf-8") as f:
            cls.C = json.load(f)
        cls.R = [r for r in cls.C["resumes"] if r["matiere"] in ("francais", "maths")]

    def test_outils_par_unite(self):
        for cours in ("CE1", "CE2", "CM1", "CM2"):
            for u in range(1, 13):
                outils = {r["domaine"] for r in self.R if r["cours"] == cours and r["matiere"] == "francais"
                          for a in r.get("ancres") or [] if a.get("type") == "unite" and a.get("unite") == u}
                self.assertTrue({"Grammaire", "Conjugaison", "Orthographe", "Vocabulaire"} <= outils, (cours, u, outils))

    def test_sons_ci_cp(self):
        import importlib.util
        for cours, mini in (("CI", 32), ("CP", 39)):
            lecons = {a["lecon"] for r in self.R if r["cours"] == cours for a in r.get("ancres") or [] if a.get("type") == "lecon"}
            self.assertGreaterEqual(len(lecons), mini, cours)

    def test_domaines_maths(self):
        doms = {"Numération et arithmétique", "Mesures", "Géométrie", "Résolution de problèmes"}
        for cours in ("CI", "CP", "CE1", "CE2", "CM1", "CM2"):
            d = {r["domaine"] for r in self.R if r["cours"] == cours and r["matiere"] == "maths"}
            self.assertEqual(d, doms, cours)
        for cours in ("CE1", "CE2", "CM1", "CM2"):
            ds = {(a["dossier"], a["sa"]) for r in self.R if r["cours"] == cours and r["matiere"] == "maths"
                  for a in r.get("ancres") or [] if a.get("type") == "dossier"}
            self.assertGreaterEqual(len(ds), 12, cours)

    def test_place_dans_annee(self):
        for r in self.R:
            self.assertIn(r.get("trimestre"), ("I", "II", "III"), r["id"])
            self.assertIn(r.get("mois"), ["Septembre", "Octobre", "Novembre", "Décembre", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin"], r["id"])
            self.assertTrue(r.get("ancres"), r["id"])
            self.assertEqual(len(r["exercices"]), 3, r["id"])
        n = [r for r in self.R if r["id"].startswith("r125-")]
        self.assertGreaterEqual(len(n), 300)
        self.assertTrue(all("à valider par le conseiller pédagogique" in r["mention"] for r in n))
        lot = self.C["lots"]["v125-resumes"]
        self.assertEqual(lot["total"], len([r for r in self.C["resumes"] if not r.get("seance_resume")]))   # v1.29 : hors résumés de séance
        brut = __import__("json").dumps(n, ensure_ascii=False)
        for interdit in ("gmail", "SONDJO", "AFROMAUD", "MARIUS", "COPAS"):
            self.assertNotIn(interdit, brut)

    def test_es_est_par_sa(self):
        """v1.27.1 : chaque SA d'Éducation sociale (dont morale et civisme) et d'EST, du CI au CM2, a ses résumés placés dans la progression."""
        lot = self.C["lots"]["v127-es-est"]
        for cours in ("CI", "CP", "CE1", "CE2", "CM1", "CM2"):
            for m in ("es", "est"):
                t = lot["table"][cours][m]
                self.assertTrue(t["sa"], (cours, m))
                self.assertEqual(t["couvertes"], t["sa"], (cours, m))
            self.assertGreater(lot["table"][cours]["es"]["emc"], 0, cours)
        rs = [r for r in self.C["resumes"] if r["matiere"] in ("es", "est") and not r.get("seance_resume")]
        self.assertEqual(len(rs), lot["resumes"])
        for r in rs:
            self.assertIn(r.get("trimestre"), ("I", "II", "III"), r["id"])
            self.assertTrue(r.get("ancres") and r.get("mois"), r["id"])
            self.assertGreaterEqual(len(r.get("vocabulaire") or []), 2, r["id"])
            self.assertTrue(r.get("je_retiens"), r["id"])
            self.assertEqual(len(r["exercices"]), 3, r["id"])
            self.assertIn("valider", r["mention"], r["id"])
            if r.get("emc"):
                self.assertTrue(r.get("situation") and r.get("engagement"), r["id"])
        par = {r["id"]: r for r in rs}
        for rid in ("res-cm1-01", "res-cm1-02", "res-cm1-03", "res-cm1-04", "res-cm2-01", "res-cm2-02", "res-cm2-03"):
            self.assertTrue(par[rid].get("sources") and par[rid].get("emc"), rid)
        self.assertIn("Romuald Wadagni", par["res-cm1-02"]["regle"] + par["res-cm1-02"]["je_retiens"])
        self.assertIn("sept ans", par["res-cm1-03"]["regle"])
        self.assertEqual([a["sa"] for a in par["r123-ce2-es-03"]["ancres"]], [3])
        sa1 = [a for a in par["res-cm1-02"]["ancres"] if a["type"] == "sa"][0]
        self.assertTrue(sa1["libelle"].startswith(("SA 1 — Le Bénin", "SA 1 — La vie démocratique")) and sa1["seances"], sa1)
        brut = __import__("json").dumps([{k: r.get(k) for k in ("vocabulaire", "je_retiens", "situation", "engagement")} for r in rs], ensure_ascii=False)
        for interdit in ("gmail", "SONDJO", "AFROMAUD", "COPAS", "PédaFiche"):
            self.assertNotIn(interdit, brut)

    def test_semaine_et_pages(self):
        r = lancer_resumes()
        self.assertIn("OK-resumes", r.stdout, r.stdout[-3000:] + r.stderr[-5000:])


SCRIPT_RES = r'''
import html, os, sys, tempfile
os.environ[os.environ["PREFIXE_SCRIPT"] + "_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
from app import app, q, ex
import m_preparation as MP
from werkzeug.security import generate_password_hash
app.config["TESTING"] = True
C = MP.corpus()
# fiches de maths CE2, dossier 1 SA 2 -> résumés rattachés au dossier 1 / SA 2
fi = [f for f in C["fiches"] if f["cours"] == "CE2" and f["matiere"] == "maths" and "DOSSIER N°1" in str(f.get("unite")) and f.get("sa") == 2]
assert fi, "fiches CE2 D1 SA2"
rs = MP.resumes_semaine("CE2", {"maths": fi})["maths"]
assert rs and all(any(a.get("dossier") == 1 and a.get("sa") == 2 for a in r["ancres"]) for r in rs if not r.get("seance_resume")), [r["titre"] for r in rs]   # v1.30 : + résumés des séances
# leçon 5 du CI : son « p »
fl = [f for f in C["fiches"] if f["cours"] == "CI" and f["matiere"] == "francais" and str(f.get("unite")) == "Leçon 5"]
assert fl
assert any("« p – P »" in r["titre"] for r in MP.resumes_semaine("CI", {"francais": fl})["francais"])
# grammaire de l'unité 6 du CE2
fg = [{"cours": "CE2", "matiere": "francais", "unite": "UNITE N°6", "sous_matiere": "Grammaire"}]
rg = MP.resumes_semaine("CE2", {"francais": fg})["francais"]
assert rg and all(r["domaine"] == "Grammaire" for r in rg if not r.get("seance_resume")), [r["titre"] for r in rg]
m = MP.resumes_mois("CM1", "Septembre")
assert m["francais"] and m["maths"]
with app.app_context():
    ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
c = app.test_client()
assert c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}).status_code == 302
t = html.unescape(c.get("/preparation/resumes").data.decode())
assert "résumés par cours et par domaine" in t and "Numération et arithmétique" in t and "Sons et graphies" in t
t = html.unescape(c.get("/preparation/resumes?cours=CE2&matiere=francais").data.decode())
assert "Trimestre I" in t and "Septembre" in t and "Unité 1" in t and "Les signes de ponctuation" in t, t[-2000:]
assert t.index("Trimestre I") < t.index("Trimestre II") < t.index("Trimestre III")
t = html.unescape(c.get("/preparation/resumes?cours=CM2&matiere=maths").data.decode())
assert "Dossier 1 · SA 1" in t
rid = next(r["id"] for r in C["resumes"] if r["id"].startswith("r125-ce1-francais"))
t = html.unescape(c.get(f"/preparation/resumes/{rid}").data.decode())
assert "Progression : trimestre" in t
# v1.27.1 : Éducation sociale (dont morale et civisme) et EST dans la progression, la semaine et le mois
fe = [f for f in C["fiches"] if f["cours"] == "CM1" and f["matiere"] == "es" and str(f.get("sa")) == "1"][:2]
assert fe
re_ = MP.resumes_semaine("CM1", {"es": fe})["es"]
assert any(r["id"] == "res-cm1-02" for r in re_), [r["id"] for r in re_]
fs = [f for f in C["fiches"] if f["cours"] == "CE2" and f["matiere"] == "est" and str(f.get("sa")) == "6"][:1]
assert any("cycle de l'eau" in r["titre"] for r in MP.resumes_semaine("CE2", {"est": fs})["est"])
assert "es" in MP.resumes_mois("CM1", "Septembre") and "est" in MP.resumes_mois("CM1", "Septembre")
t = html.unescape(c.get("/preparation/resumes?cours=CM1&matiere=es").data.decode())
assert "Trimestre I" in t and ("SA 1 — Le Bénin" in t or "SA 1 — La vie démocratique" in t) and "SA couvertes par cours" in t, t[-2500:]
t = html.unescape(c.get("/preparation/resumes?cours=CP&matiere=emc").data.decode())
assert "Je sais saluer" in t and "activités des gens" not in t
t = html.unescape(c.get("/preparation/resumes/r123-cp-es-01").data.decode())
assert "Situation de la vie courante" in t and "Je m'engage" in t and "Valeur ou règle de vie" in t and "Vocabulaire" in t
t = html.unescape(c.get("/preparation/resumes/res-cm2-02").data.decode())
assert "Sources :" in t and "petite évaluation" in t
r = c.get("/preparation/resumes/res-cm2-02.pdf")
assert r.status_code == 200 and r.mimetype == "application/pdf"
print("OK-resumes")
'''


def lancer_resumes():
    env = dict(os.environ)
    for k in ("SCOLARIS_EDITION", "SCOLARIS_DATA", "GESCOL_EDITION", "GESCOL_DATA"):
        env.pop(k, None)
    env["PREFIXE_SCRIPT"] = PREFIXE
    env[PREFIXE + "_EDITION"] = "publique"
    return subprocess.run([sys.executable, "-c", SCRIPT_RES], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)


class TestIconesCommunes(unittest.TestCase):
    def test_meme_jeu_que_les_applications_js(self):
        import re
        sys.path.insert(0, RACINE)
        import icones
        self.assertGreaterEqual(len(icones.NOMS), 133)
        self.assertEqual(icones.NOMS[:3], ["accueil", "ajouter", "fermer"])
        self.assertEqual(icones.NOMS[132], "examen")
        for n in ("ballon", "stade", "examen", "patrimoine", "inspection", "infirmerie"):
            self.assertIn(f"ic-{n}", icones.svg(n))
            self.assertIn(f"ic-{n}", icones.svg(icones.code(n)))
        self.assertEqual(icones.retirer_codes("A" + icones.code("ballon") + "B"), "AB")
        self.assertFalse(re.search(r"[\U0001F000-\U0001FAFF]", "".join(icones.D.values())))


if __name__ == "__main__":
    unittest.main()
