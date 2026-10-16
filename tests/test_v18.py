# -*- coding: utf-8 -*-
"""Tests de la version 1.8 : « Ma classe » et accueil par rôle, appel rapide et saisies hors ligne, groupes de
remédiation, alertes de pilotage (CS / DDEMP), bandeau de mise à jour et file d'attente hors ligne.

Les parcours complets tournent dans des sous-processus (une édition par processus).
Lancer : python3 -m unittest tests.test_v18"""
import os
import subprocess
import sys
import unittest
from datetime import date

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)

import a_alertes as AL  # noqa: E402
import m_classe as MC  # noqa: E402


def lancer(script, edition):
    env = dict(os.environ)
    env.pop("SCOLARIS_EDITION", None)
    env.pop("SCOLARIS_DATA", None)
    env["SCOLARIS_EDITION_SCRIPT"] = edition
    r = subprocess.run([sys.executable, "-c", script], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)
    return r


class TestRemediation(unittest.TestCase):
    def test_groupes_par_note(self):
        self.assertEqual(MC.groupe_de_note(0), "g1")
        self.assertEqual(MC.groupe_de_note(6.99), "g1")
        self.assertEqual(MC.groupe_de_note(7), "g2")
        self.assertEqual(MC.groupe_de_note(10), "g3")
        self.assertEqual(MC.groupe_de_note(15), "g4")
        self.assertEqual(MC.groupe_de_note(20), "g4")
        self.assertIsNone(MC.groupe_de_note(None))

    def test_activites(self):
        self.assertIn("Calcul mental", MC.activites_pour("Mathématiques"))
        self.assertIn("Lecture", MC.activites_pour("Lecture"))
        self.assertIn("Dictées", MC.activites_pour("Dictée"))
        self.assertEqual(MC.activites_pour("Matière inconnue"), MC.ACTIVITE_DEFAUT)

    def test_groupes_remediation(self):
        mats = [{"id": 1, "nom": "Mathématiques"}, {"id": 2, "nom": "Lecture"}]
        el = [{"id": i, "nom": f"E{i}", "prenoms": ""} for i in range(5)]
        notes = [{1: 3, 2: 12}, {1: 8, 2: 16}, {1: 11, 2: None}, {1: 17, 2: 9}, {1: None, 2: None}]
        lignes = []
        for e, n in zip(el, notes):
            vals = [v for v in n.values() if v is not None]
            lignes.append({"eleve": e, "notes": n, "dispenses": set(), "appreciations": {}, "compose": bool(vals),
                           "maitrise": (sum(1 for v in vals if v >= 10) >= 2) if vals else None,
                           "moyenne": sum(vals) / len(vals) if vals else None,
                           "reussies": sum(1 for v in vals if v >= 10), "seuil": 2})
        r = MC.groupes_remediation(mats, lignes)
        m = r["matieres"][0]
        self.assertEqual([e["id"] for e in m["groupes"]["g1"]], [0])
        self.assertEqual([e["id"] for e in m["groupes"]["g2"]], [1])
        self.assertEqual([e["id"] for e in m["groupes"]["g3"]], [2])
        self.assertEqual([e["id"] for e in m["groupes"]["g4"]], [3])
        self.assertEqual([e["id"] for e in m["absents"]], [4])
        self.assertEqual(m["taux"], 50.0)
        self.assertEqual(r["matieres"][1]["taux"], 66.7)
        self.assertEqual(r["a_renforcer"], [])
        nm = {x["eleve"]["id"]: x for x in r["non_maitrise"]}
        self.assertEqual(set(nm), {0, 1, 2, 3})
        self.assertEqual(nm[0]["matieres"], ["Mathématiques"])
        self.assertEqual(nm[3]["matieres"], ["Lecture"])

    def test_maternelle(self):
        mats = [{"id": 1, "nom": "Langage"}]
        lignes = [{"eleve": {"id": i, "nom": f"E{i}", "prenoms": ""}, "notes": {1: None}, "dispenses": set(),
                   "appreciations": {1: a}, "compose": True, "maitrise": a != "NA", "moyenne": None, "reussies": 0,
                   "seuil": 1} for i, a in enumerate(["NA", "ECA", "A"])]
        g = MC.groupes_remediation(mats, lignes)["matieres"][0]["groupes"]
        self.assertEqual([len(g[k]) for k in ("g1", "g2", "g3", "g4")], [1, 1, 1, 0])


class TestAlertesRegles(unittest.TestCase):
    def test_periode(self):
        self.assertEqual(AL.periode_du_jour(date(2026, 10, 16)), ("2026-2027", "T1"))
        self.assertEqual(AL.periode_du_jour(date(2027, 2, 1)), ("2026-2027", "T2"))
        self.assertEqual(AL.periode_du_jour(date(2027, 5, 1)), ("2026-2027", "T3"))

    def test_avancement(self):
        self.assertEqual(AL.avancement("2026-2027", "T1", date(2026, 8, 1)), 0.0)
        self.assertEqual(AL.avancement("2026-2027", "T1", date(2027, 1, 5)), 1.0)
        self.assertAlmostEqual(AL.avancement("2026-2027", "T1", date(2026, 10, 31)), 60 / 121, places=3)
        self.assertEqual(AL.avancement("2026-2027", "T3", date(2027, 7, 15)), 1.0)

    def test_ecart(self):
        self.assertIsNone(AL.ecart_programmation(0, 0, .5))
        self.assertIsNone(AL.ecart_programmation(1, 0, .5))        # rien d'attendu encore
        self.assertEqual(AL.ecart_programmation(20, 2, .5), 40.0)    # 10 attendues, 2 faites
        self.assertEqual(AL.ecart_programmation(20, 12, .5), -10.0)  # en avance


SCRIPT_ECOLE = r'''
import html, io, os, re, sys, tempfile, subprocess, json
ED = os.environ["SCOLARIS_EDITION_SCRIPT"]
D = tempfile.mkdtemp()
os.environ["SCOLARIS_EDITION"] = ED; os.environ["SCOLARIS_DATA"] = D
sys.path.insert(0, os.getcwd())
subprocess.run([sys.executable, "demo.py"], check=True, env=dict(os.environ), stdout=subprocess.DEVNULL)
import config
from app import app, q, ex
app.config["TESTING"] = True
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.data[-800:]); return r
def txt(r): return html.unescape(r.data.decode())
def client(ident, mdp="demo1234"):
    c = app.test_client(); ok(c.post("/login", data={"identifiant": ident, "mdp": mdp}), 302); return c
with app.app_context():
    ex("UPDATE utilisateurs SET doit_changer_mdp=0 WHERE identifiant='admin'")
    cm2 = q("SELECT id FROM classes WHERE niveau='CM2' ORDER BY id", un=True)["id"]
    autre = q("SELECT id FROM classes WHERE id<>? ORDER BY id", (cm2,), un=True)["id"]
    uid = q("SELECT id FROM utilisateurs WHERE identifiant='enseignant'", un=True)["id"]
    ex("INSERT INTO affectations (user_id, classe_id, matiere_id) VALUES (?,?,NULL)", (uid, cm2))
    ev = ex("INSERT INTO evaluations (classe_id, annee, type, numero, libelle, date) VALUES (?,?,?,?,?,?)",
            (cm2, config.ANNEE_ACTIVE, "sommative", 1, config.PERIODES[0], "2000-01-01"))
    ev = q("SELECT id FROM evaluations WHERE classe_id=? ORDER BY id DESC", (cm2,), un=True)["id"]
    eleves = [r["id"] for r in q("SELECT id FROM eleves WHERE classe_id=? AND statut='inscrit'", (cm2,))]
    from datetime import date, timedelta
    for k in range(3):
        ex("INSERT OR IGNORE INTO presences (eleve_id, date, statut, justifie) VALUES (?,?,?,0)",
           (eleves[0], (date.today() - timedelta(days=k + 1)).isoformat(), "absent"))

# enseignant : « / » → Ma classe → sa classe
c = client("enseignant")
r = ok(c.get("/"), 302); assert r.headers["Location"].endswith("/ma-classe"), r.headers["Location"]
r = ok(c.get("/ma-classe"), 302); assert r.headers["Location"].endswith(f"/ma-classe/{cm2}")
t = txt(ok(c.get(f"/ma-classe/{cm2}")))
for k in ("Faire l'appel", "Groupes de remédiation", "Élèves à suivre", "absence(s) non justifiée(s)", "Évaluation sommative d'étape N° 1"):
    assert k in t, k
assert c.get(f"/ma-classe/{autre}").status_code in (403, 404)
# appel rapide : formulaire gardé hors ligne, compteurs
t = txt(ok(c.get(f"/appel/{cm2}/{date.today().isoformat()}")))
assert "data-hors-ligne" in t and 'id="tous-presents"' in t and 'id="n-a"' in t, t[:300]
# saisie des notes : hors ligne + bouton remédiation
t = txt(ok(c.get(f"/evaluations/{ev}/saisie")))
assert "data-hors-ligne" in t and f"/evaluations/{ev}/remediation" in t
# remédiation
t = txt(ok(c.get(f"/evaluations/{ev}/remediation")))
assert "Groupe 1 – Soutien intensif" in t or "Groupe 3 – Maîtrise" in t, t[-1500:]
assert "Élèves en non-maîtrise" in t
assert ok(c.get(f"/evaluations/{ev}/remediation.pdf")).data.startswith(b"%PDF")
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(ok(c.get(f"/evaluations/{ev}/remediation.xlsx")).data))
assert wb.active.max_row > 3
# une évaluation d'une autre classe : refusée
with app.app_context():
    ex("INSERT INTO evaluations (classe_id, annee, type, numero, libelle) VALUES (?,?,?,?,?)",
       (autre, config.ANNEE_ACTIVE, "formative", 1, "Formative 1"))
    ev2 = q("SELECT id FROM evaluations WHERE classe_id=?", (autre,), un=True)["id"]
assert c.get(f"/evaluations/{ev2}/remediation").status_code in (403, 404)
# menu : Ma classe
assert "Ma classe" in txt(c.get(f"/ma-classe/{cm2}"))
# directeur : raccourcis et état des sauvegardes sur le tableau de bord
d = client("directeur")
t = txt(ok(d.get("/")))
assert 'class="raccourcis"' in t and "Faire l'appel" in t, t[:500]
a = client("admin", "admin123")
t = txt(ok(a.get("/")))
assert "Sauvegardes" in t and ("Aucune sauvegarde" in t or "Dernière" in t)
# bandeau de mise à jour (application Windows) et file d'attente hors ligne
t = txt(ok(a.get("/", headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})))
assert 'id="file-attente"' in t
assert 'id="bandeau-maj"' in t and 'data-depot="' + config.DEPOT_MAJ + '"' in t and ".exe" in t, "bandeau"
# service worker : pages de classe gardées hors ligne
sw = ok(a.get("/sw.js")).data.decode()
assert "ma-classe" in sw
print("ECOLE-OK")
'''

SCRIPT_CS = r'''
import html, os, sys, tempfile
os.environ["SCOLARIS_EDITION"] = os.environ["SCOLARIS_EDITION_SCRIPT"]; os.environ["SCOLARIS_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config, demo_juridiction
from app import app, q, ex
app.config["TESTING"] = True
demo_juridiction.main()
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.data[-800:]); return r
c = app.test_client()
ok(c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"}), 302)
t = html.unescape(ok(c.get("/juridiction/alertes")).data.decode())
assert "Alertes de pilotage" in t and "Urgentes" in t
with app.app_context():
    n = q("SELECT COUNT(*) n FROM jur_ecoles", un=True)["n"]
    # une école sans transmission → alerte haute
    tr = q("SELECT COUNT(DISTINCT ecole_id) n FROM jur_sections WHERE annee=?", (config.ANNEE_ACTIVE,), un=True)["n"]
if tr < n:
    assert "sans aucune transmission" in t, t[-2000:]
t = html.unescape(ok(c.get("/juridiction/tableau-de-bord")).data.decode())
assert "Alertes de pilotage" in t
assert 'class="raccourcis' in t
print("CS-OK")
'''


class TestParcoursEcole(unittest.TestCase):
    def test_publique(self):
        r = lancer(SCRIPT_ECOLE, "publique")
        self.assertIn("ECOLE-OK", r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_privee(self):
        r = lancer(SCRIPT_ECOLE, "privee")
        self.assertIn("ECOLE-OK", r.stdout, r.stdout[-3000:] + r.stderr[-6000:])


class TestAlertesJuridiction(unittest.TestCase):
    def test_cs(self):
        r = lancer(SCRIPT_CS, "cs")
        self.assertIn("CS-OK", r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_ddemp(self):
        r = lancer(SCRIPT_CS, "ddemp")
        self.assertIn("CS-OK", r.stdout, r.stdout[-3000:] + r.stderr[-6000:])


class TestScripts(unittest.TestCase):
    def test_js_file_et_maj(self):
        with open(os.path.join(RACINE, "static", "scolaris.js"), encoding="utf-8") as f:
            js = f.read()
        for k in ("scolaris-file-attente", "function viderFile", "function initMaj", "releases/latest", "data-hors-ligne"):
            self.assertIn(k, js)


if __name__ == "__main__":
    unittest.main()
