# -*- coding: utf-8 -*-
"""Tests de la version 1.7 :
* référentiel des communes, arrondissements et villages (5 290 villages et quartiers) ;
* en-tête officiel (armoiries, emblème du MEMP) sur les fiches et cartes ; signature digitale du parent (fiche CEP) ;
* module « Conseillers pédagogiques » (a_cp) : calculs du bilan trimestriel, programmation, journal, bilan
  (Enregistrer, Charger, Télécharger PDF / Excel, Recalculer, Réinitialiser, Supprimer), envoi à la DDEMP.
Lancer : python3 -m unittest tests.test_v17   (depuis la racine du projet)

La partie CS est jouée dans un sous-processus (édition « cs », base de démonstration demo_juridiction)."""
import base64
import html as html_mod
import io
import os
import re
import subprocess
import sys
import tempfile
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RACINE)
EDITION = os.environ.get("SCOLARIS_EDITION_TEST", "publique")
os.environ.setdefault("SCOLARIS_EDITION", EDITION)
os.environ.setdefault("SCOLARIS_DATA", tempfile.mkdtemp(prefix=f"scolaris_v17_{EDITION}_"))

from app import app  # noqa: E402
import config  # noqa: E402
import a_cp  # noqa: E402
import referentiel_villages as RV  # noqa: E402

EN_PLACE = config.EDITION == EDITION


def txt(r):
    return html_mod.unescape(r.data.decode())


def png_signature():
    from PIL import Image, ImageDraw
    img = Image.new("RGBA", (300, 120), (0, 0, 0, 0))
    ImageDraw.Draw(img).line([(10, 100), (80, 20), (150, 100), (220, 30), (290, 90)], fill=(20, 40, 160, 255), width=6)
    b = io.BytesIO()
    img.save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


class TestReferentielVillages(unittest.TestCase):
    def test_totaux(self):
        self.assertEqual(RV.totaux(), {"departements": 12, "communes": 77, "arrondissements": 546, "villages": 5290})

    def test_cascade(self):
        self.assertEqual(RV.communes("PLATEAU"), ["Adja-Ouèrè", "Ifangni", "Sakété", "Kétou", "Pobè"])
        self.assertEqual(RV.departement_de("pobe"), "Plateau")
        self.assertIn("Issaba", RV.arrondissements("Pobè"))
        self.assertIn("Igbo-Ewé", RV.villages("Pobè", "Issaba"))
        self.assertEqual(len(RV.arrondissements("Cotonou")), 13)
        self.assertEqual(RV.nom_officiel_commune("SEME PODJI"), "Sèmè-Podji")


class TestCalculsBilan(unittest.TestCase):
    """Chiffres du canevas (CS de Pobè, trimestre 1) : totaux et pourcentages."""

    def bilan(self):
        act = [{"cle": "ccs", "lib": "C/CS"}, {"cle": "z1", "lib": "CP Zone 1"}, {"cle": "z2", "lib": "CP Zone 2"}]
        b = {"acteurs": act, "s1": {}, "s2": {}, "s3": {}, "s4": {}, "s5": {"lignes": {}}}
        b["s1"] = {"ccs": dict(em_ecoles=32, em_sections=62, ep_ecoles=125, ep_classes=599, ens_em_h=21, ens_em_f=46,
                               ens_ep_h=394, ens_ep_f=209, up_ep=25, up_em=3),
                   "z1": dict(em_ecoles=11, em_sections=20, ep_ecoles=44, ep_classes=228, ens_em_h=3, ens_em_f=27,
                              ens_ep_h=121, ens_ep_f=140, up_ep=13, up_em=1),
                   "z2": dict(em_ecoles=11, em_sections=21, ep_ecoles=41, ep_classes=203, ens_em_h=14, ens_em_f=10,
                              ens_ep_h=130, ens_ep_f=43, up_ep=6, up_em=1)}
        b["s2"] = {"ccs": dict(vc_prog_em=2, vc_prog_t=15, vc_real_em=1, vc_real_t=23, up_prog=3, up_prep=3, up_real=3,
                               fo_prog=2, fo_real=2),
                   "z1": dict(vc_prog_em=5, vc_prog_t=45, vc_real_em=5, vc_real_t=50, up_prog=3, up_prep=3, up_real=3,
                              fo_prog=2, fo_real=2),
                   "z2": dict(vc_prog_em=7, vc_prog_t=45, vc_real_em=10, vc_real_t=62, up_prog=3, up_prep=3, up_real=3,
                              fo_prog=2, fo_real=2)}
        b["s4"] = {k: dict(vc_cible_em=b["s1"][k]["em_sections"], vc_cible_t=b["s1"][k]["em_sections"] + b["s1"][k]["ep_classes"],
                           fo_benef_em=50, fo_benef_t=314, fo_cible_em=67, fo_cible_t=670) for k in ("ccs", "z1", "z2")}
        b["s5"]["lignes"] = {"CP": dict(ins_g=3026, ins_f=2662, comp_g=2658, comp_f=2316, min_g=2051, min_f=1872)}
        return a_cp.totaux(b)

    def test_reseau(self):
        b = self.bilan()
        t = b["s1"]["_tcs"]
        self.assertEqual((t["t_ecoles"], t["t_classes"], t["ens_t"], t["up_t"]), (157, 661, 670, 28))  # = ligne C/CS
        self.assertEqual(b["s1"]["_tcp"]["em_ecoles"], 22)

    def test_activites(self):
        b = self.bilan()
        self.assertEqual(b["s2"]["ccs"]["vc_pct_t"], 153.33)
        self.assertEqual(b["s2"]["z1"]["vc_pct_em"], 100.0)
        self.assertEqual(b["s2"]["_tcs"]["vc_real_t"], 135)
        # séances d'UP et formations : activités communes, le total n'est pas une somme
        self.assertEqual((b["s2"]["_tcs"]["up_real"], b["s2"]["_tcs"]["fo_real"]), (3, 2))

    def test_couverture(self):
        b = self.bilan()
        self.assertEqual(b["s4"]["ccs"]["vc_cov_t"], 3.48)              # 23 / 661
        self.assertEqual(b["s4"]["ccs"]["vc_cov_em"], 1.61)             # 1 / 62
        self.assertEqual(b["s4"]["_tcs"]["fo_cov_t"], 46.87)            # 314 / 670 (formation commune comptée une fois)

    def test_evaluation_et_analyse(self):
        b = self.bilan()
        r = b["s5"]["lignes"]["CP"]
        self.assertEqual((r["reu_g"], r["reu_f"], r["reu_t"]), (77.16, 80.83, 78.87))
        an = a_cp.analyse_auto(b)
        self.assertIn("5688 apprenants", an["quantitative"])
        self.assertIn("les filles sont plus performant", an["genre"])

    def test_outils(self):
        self.assertEqual(a_cp.bornes("2026-2027", "T2"), ("2027-01-01", "2027-03-31"))
        self.assertEqual(a_cp.annee_de("2026-10-05"), "2026-2027")
        self.assertEqual(a_cp.annee_de("2027-02-01"), "2026-2027")
        self.assertEqual(a_cp.entier("1 234"), 1234)
        self.assertEqual(a_cp.pct_txt(66.666), "66,67 %")
        self.assertEqual(a_cp.lib_zone("2"), "CP Zone 2")


class TestEnteteEtSignature(unittest.TestCase):
    """Édition de l'école : emblème du MEMP / armoiries sur les documents, signature du parent sur la fiche CEP."""

    @classmethod
    def setUpClass(cls):
        if not EN_PLACE or config.juridiction():
            raise unittest.SkipTest("édition d'école requise")
        app.config["TESTING"] = True
        from werkzeug.security import generate_password_hash
        with app.app_context():
            from app import ex
            ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'",
               (generate_password_hash("admin123"),))
            cl = ex("INSERT INTO classes (nom, niveau, frais_scolarite, ecole_id) VALUES (?,?,?,?)", ("CM2 V17", "CM2", 0, 1))
            cls.el = ex("""INSERT INTO eleves (nom, prenoms, sexe, date_naissance, lieu_naissance, classe_id, statut,
                           date_inscription, annee, matricule, parent_nom) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                        ("HOUNSA", "Inès", "F", "2015-05-03", "Igbo-Ewé (Pobè)", cl, "inscrit", "2026-09-15",
                         config.ANNEE_ACTIVE, "V17-0001", "HOUNSA Paul"))
        cls.c = app.test_client()
        assert cls.c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}).status_code == 302

    def test_images_officielles(self):
        import pdfs
        self.assertIsNotNone(pdfs.image_officielle("armoiries"))
        self.assertIsNotNone(pdfs.image_officielle("memp"))

    def test_documents_et_cartes(self):
        for u in (f"/eleves/{self.el}/document/carte.pdf", f"/eleves/{self.el}/document/certificat_frequentation.pdf",
                  f"/cep/fiche/{self.el}/pdf", f"/carte/{self.el}/pdf"):
            r = self.c.get(u)
            self.assertEqual(r.status_code, 200, u)
            self.assertTrue(r.data.startswith(b"%PDF"), u)
            self.assertIn(b"/Image", r.data, u)               # armoiries / emblème du MEMP dessinés

    def test_signature_parent(self):
        r = self.c.post(f"/signatures/parent:{self.el}", json={"image": png_signature()})
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.get_json()["ok"])
        self.assertIn(f"parent:{self.el}", txt(self.c.get(f"/cep/fiche/{self.el}")))
        sans = len(self.c.get(f"/eleves/{self.el}/document/fiche_cep.pdf").data)
        self.c.post(f"/signatures/parent:{self.el}", json={"effacer": 1})
        self.assertLess(len(self.c.get(f"/eleves/{self.el}/document/fiche_cep.pdf").data), sans)

    def test_champs_geo(self):
        t = txt(self.c.get("/eleves/nouveau")) if self.c.get("/eleves/nouveau").status_code == 200 else ""
        if t:
            self.assertIn('data-geo="lieu"', t)
        self.assertEqual(self.c.get("/static/referentiel_villages.json").status_code, 200)


SCRIPT_CS = r'''
import io, json, os, re, sys, html, tempfile
os.environ["SCOLARIS_EDITION"] = "cs"; os.environ["SCOLARIS_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config, demo_juridiction
from app import app
app.config["TESTING"] = True
demo_juridiction.main()
c = app.test_client()
assert c.post("/login", data={"identifiant": "responsable", "mdp": "demo1234"}).status_code == 302
A = config.ANNEE_ACTIVE; a1 = A[:4]
with app.app_context():
    from app import q
    z = q("SELECT * FROM jur_zones ORDER BY ordre", un=True)["id"]
def ok(r, code=200):
    assert r.status_code == code, (r.status_code, r.data[-600:]); return r
ok(c.get("/juridiction/conseillers"))
ok(c.post(f"/juridiction/conseillers/{z}", data={"cp_nom": "ATCHATIN Salomon", "grade": "CP"}), 302)
ok(c.post("/juridiction/conseillers/programmation", data={"annee": A, "periode": "T1", "p_ccs_vc_t": "15", f"p_z{z}_vc_t": "45",
                                                            "p_ccs_co": "3"}), 302)
x = ok(c.get("/juridiction/conseillers/programmation.xlsx", query_string={"annee": A, "periode": "T1"})).data
ok(c.post("/juridiction/conseillers/programmation/import", data={"annee": A, "periode": "T2", "fichier": (io.BytesIO(x), "p.xlsx")},
          content_type="multipart/form-data"), 302)
for d in (dict(categorie="suivi", acteur=f"z{z}", cycle="EM", date_debut=f"{a1}-10-12", nb="2"),
          dict(categorie="concertation", acteur="commune", date_debut=f"{a1}-10-05", designation="Concertation"),
          dict(categorie="autre", acteur="ccs", date_debut=f"{a1}-09-16", lieu="Écoles", designation="Tournée de rentrée")):
    ok(c.post("/juridiction/conseillers/activites/nouvelle", data={"statut": "realisee", **d}), 302)
ok(c.get("/juridiction/conseillers/activites", query_string={"annee": A, "periode": "T1"}))
ok(c.get("/juridiction/conseillers/activites.xlsx", query_string={"annee": A, "periode": "T1"}))
t = html.unescape(ok(c.get("/juridiction/conseillers/bilan", query_string={"annee": A, "periode": "T1"})).data.decode())
vals = dict(re.findall(r'name="(s[1-4]_[a-z0-9_]+)" value="([^"]*)"', t))
assert vals["s2_ccs_vc_prog_t"] == "15" and vals[f"s2_z{z}_su_real_em"] == "2" and vals["s3_ccs_co_real"] == "1", vals
assert "Tournée de rentrée" in t
ok(c.post("/juridiction/conseillers/bilan", data={"annee": A, "periode": "T1", "action": "enregistrer", **vals,
                                                  "commentaires": "Efforts significatifs", "au_designation": ["Tournée de rentrée"],
                                                  "au_lieu": ["Écoles"], "au_periode": ["16/09"], "au_participants": ["C/CS"]}), 302)
pdf = ok(c.get("/juridiction/conseillers/bilan.pdf", query_string={"annee": A, "periode": "T1"})).data
assert pdf.startswith(b"%PDF")
x = ok(c.get("/juridiction/conseillers/bilan.xlsx", query_string={"annee": A, "periode": "T1"})).data
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(x))
assert {"1-Reseau", "2-Activites", "4-Couverture", "5-Evaluation", "7-Textes"} <= set(wb.sheetnames)
for row in wb["7-Textes"].iter_rows():
    if row[0].value == "commentaires":
        row[2].value = "Chargé depuis Excel"
b = io.BytesIO(); wb.save(b); b.seek(0)
ok(c.post("/juridiction/conseillers/bilan/import", data={"annee": A, "periode": "T1", "fichier": (b, "b.xlsx")},
          content_type="multipart/form-data"), 302)
assert "Chargé depuis Excel" in html.unescape(c.get("/juridiction/conseillers/bilan", query_string={"annee": A, "periode": "T1"}).data.decode())
ok(c.post("/juridiction/conseillers/bilan", data={"annee": A, "periode": "T1", "action": "recalculer"}), 302)
ok(c.get("/juridiction/conseillers/bilan.pdf", query_string={"annee": A, "periode": "T1", "zone": z}))
ok(c.post("/juridiction/conseillers/bilan", data={"annee": A, "periode": "AN", "action": "valider"}), 302)
with app.app_context():
    from app import q
    sec = json.loads(q("SELECT donnees FROM jur_cs_sections WHERE section='cp_bilans'", un=True)["donnees"])
    assert len(sec) == 2 and all("acteurs" in s["bilan"] for s in sec), sec
ok(c.post("/juridiction/conseillers/bilan", data={"annee": A, "periode": "AN", "action": "supprimer"}), 302)
ok(c.post("/juridiction/conseillers/bilan", data={"annee": A, "periode": "T1", "action": "reinitialiser"}), 302)
with app.app_context():
    from app import q
    assert q("SELECT COUNT(*) n FROM jur_cp_bilans", un=True)["n"] == 0
print("CS-OK")
'''


class TestModuleCP(unittest.TestCase):
    def test_parcours_cs(self):
        env = dict(os.environ)
        env.pop("SCOLARIS_EDITION", None)
        env.pop("SCOLARIS_DATA", None)
        r = subprocess.run([sys.executable, "-c", SCRIPT_CS], cwd=RACINE, env=env, capture_output=True, text=True, timeout=900)
        self.assertIn("CS-OK", r.stdout, r.stderr[-4000:])


if __name__ == "__main__":
    unittest.main()
