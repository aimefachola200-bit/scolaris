# -*- coding: utf-8 -*-
"""Tests 1.22 : bureaux de l'UASCEP — membres et responsables par école, UP, zone, CS et DDEMP (onglets, ajouter,
modifier, supprimer, réinitialiser, renouvellement, anciens bureaux, synthèse, PDF, Excel commun export / import) ;
éditions école privée / publique (bureau de l'école), CS (écoles, UP, zones, CS), DDEMP (tous niveaux) ; enseignant
en lecture seule.
Lancer : python3 -m unittest tests.test_v122"""
import os
import subprocess
import sys
import unittest

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PREFIXE = "GESCOL" if os.path.exists(os.path.join(RACINE, "GESCOL_PUBLIC.md")) else "SCOLARIS"

SCRIPT = r'''
import html, io, json, os, sys, tempfile
P = os.environ["PREFIXE_SCRIPT"]
ED = os.environ["EDITION_SCRIPT"]
os.environ[P + "_EDITION"] = ED; os.environ[P + "_DATA"] = tempfile.mkdtemp()
sys.path.insert(0, os.getcwd())
import config
from app import app, q, ex
from werkzeug.security import generate_password_hash
app.config["TESTING"] = True
def ok(r, code=200):
    assert r.status_code == code, (r.request.path, r.status_code, r.headers.get("Location"), r.data[-1500:]); return r
def txt(r): return html.unescape(r.data.decode())
import m_uascep as U
assert U.fin_mandat("2026-10-15", 2) == "2028-10-15" and U.statut_mandat("2020-01-01") == "Expiré" and U.statut_mandat("") == "Date de fin non renseignée"
assert U.niveau_de("Zone pédagogique") == "ZONE" and U.niveau_de("ECOLE") == "ECOLE" and U.niveau_de("Département (DDEMP)") == "DDEMP"
JUR = ED in ("cs", "ddemp")
with app.app_context():
    ex("UPDATE utilisateurs SET mdp_hash=?, doit_changer_mdp=0 WHERE identifiant='admin'", (generate_password_hash("admin123"),))
    if JUR:
        ex("INSERT OR REPLACE INTO jur_structure (id, niveau, departement, cs, nom) VALUES (1, ?, 'Plateau', 'Pobè', 'CS Pobè')", (ED,))
        cs = ex("INSERT INTO jur_cs (departement, nom) VALUES ('Plateau', 'Pobè')")
        z1 = ex("INSERT INTO jur_zones (cs_id, nom, ordre) VALUES (?, 'Zone I', 1)", (cs,))
        e1 = ex("INSERT INTO jur_ecoles (code, nom, cs_id, zone_id) VALUES ('EPP01', 'EPP Towé', ?, ?)", (cs, z1))
        ex("INSERT INTO jur_up (cs_id, zone_id, nom, rup_nom) VALUES (?, ?, 'UP 3', 'ADJAÏ Rachidi')", (cs, z1))
    else:
        uid = ex("INSERT INTO utilisateurs (identifiant, nom, role, mdp_hash, doit_changer_mdp, actif, ecole_id) VALUES ('ens1','HOUNSOU Awa','enseignant',?,0,1,1)",
                 (generate_password_hash("ens-mdp-123"),))
c = app.test_client()
ok(c.post("/login", data={"identifiant": "admin", "mdp": "admin123"}), 302)
t = txt(ok(c.get("/uascep")))
assert "Bureaux de l'UASCEP" in t and "Anciens" in t, t[:3000]
assert "Synthèse" in t and "Anciens bureaux" in t
assert ("Circonscription scolaire" in t) == JUR and ("Département (DDEMP)" in t) == (ED == "ddemp")
ent = str(e1) if JUR else "1"
if JUR:
    assert "EPP Towé" in t, t[t.find("<main"):][:4000]
# ajouter / modifier / supprimer
ok(c.post(f"/uascep/membre/0?niveau=ECOLE&entite={ent}", data={"ordre": "1", "nom": "AKPO", "prenoms": "Mariam", "sexe": "F", "fonction": "Présidente",
                                                              "qualite": "Enseignante", "date_election": "2026-10-15", "duree": "2", "ecole": "EPP Towé", "up": "UP 3", "zone": "Zone I"}), 302)
ok(c.post(f"/uascep/membre/0?niveau=ECOLE&entite={ent}", data={"ordre": "2", "nom": "DOSSOU", "prenoms": "Koffi", "sexe": "M", "fonction": "Secrétaire général", "date_election": "2026-10-15"}), 302)
ok(c.post(f"/uascep/membre/0?niveau=ECOLE&entite={ent}", data={"ordre": "3", "nom": "", "prenoms": "X"}), 400)
t = txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
assert "AKPO Mariam" in t and "15/10/2028" in t and "En cours" in t and "DOSSOU Koffi" in t
with app.app_context():
    mid = q("SELECT id FROM uascep_membres WHERE nom='DOSSOU'", un=True)["id"]
    bid = q("SELECT bureau_id FROM uascep_membres WHERE id=?", (mid,), un=True)["bureau_id"]
ok(c.post(f"/uascep/membre/{mid}", data={"ordre": "2", "nom": "DOSSOU", "prenoms": "Koffi Jean", "sexe": "M", "fonction": "Trésorier général", "date_election": "2026-10-15"}), 302)
assert "Koffi Jean" in txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
ok(c.post(f"/uascep/bureau/{bid}", data={"entite_lib": "EPP Towé", "date_election": "2026-10-15", "duree": "2", "ag_date": "2026-10-10", "ag_lieu": "Salle du CM2"}), 302)
# PDF et Excel (spécification commune)
r = ok(c.get(f"/uascep.pdf?niveau=ECOLE&entite={ent}")); assert r.mimetype == "application/pdf"
r = ok(c.get("/uascep.xlsx")); assert r.data[:2] == b"PK"
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(r.data))
assert wb.sheetnames == ["Membres UASCEP", "Bureaux"]
ws = wb["Membres UASCEP"]
ent_xl = [c.value for c in next(ws.iter_rows(min_row=1, max_row=6)) if c.value]
lignes = [[c.value for c in r_] for r_ in ws.iter_rows()]
k = next(i for i, l in enumerate(lignes) if l and l[0] == "Niveau")
assert lignes[k][:20] == U.ENT_MEMBRES, lignes[k]
assert any(l[3] == "AKPO" and l[14] == "15/10/2026" for l in lignes[k + 1:])
# renouvellement puis anciens bureaux
ok(c.post(f"/uascep/bureau/{bid}/renouveler", data={"date_election": "2028-10-20"}), 302)
t = txt(ok(c.get("/uascep?onglet=ANCIENS")))
assert "AKPO" in t and "archivé" in t
t = txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
assert "20/10/2030" in t and "AKPO Mariam" in t
# synthèse
t = txt(ok(c.get("/uascep?onglet=SYNTHESE")))
assert "Femmes : 1" in t and "Présidente" in t
# réinitialiser (archive) : bureau vidé
with app.app_context():
    bid2 = q("SELECT id FROM uascep_bureaux WHERE statut='actif' AND niveau='ECOLE'", un=True)["id"]
ok(c.post(f"/uascep/bureau/{bid2}/reinitialiser", data={"archive": "1"}), 302)
assert "AKPO Mariam" not in txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
# import du fichier exporté (rechargement) : membres recréés
buf = io.BytesIO()
wb2 = openpyxl.Workbook(); w = wb2.active; w.title = "Membres UASCEP"; w.append(U.ENT_MEMBRES)
w.append(["École", "EPP Towé" if JUR else config.ECOLE.get("nom", ""), 1, "BIO", "Salimatou", "F", "EPP Towé", "UP 3", "Zone I", "", "Présidente", "Directrice", "", "", "01/11/2026", 2, "", "", "", ""])
w.append(["Inconnu", "X", 1, "NOM", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])
if JUR:
    w.append(["UP", "UP 3", 1, "KPADE", "Jean", "M", "", "UP 3", "Zone I", "ADJAÏ Rachidi", "Président", "Directeur", "", "", "05/11/2026", 2, "", "", "", ""])
    w.append(["Zone pédagogique", "Zone I", 1, "SOSSA", "Ida", "F", "", "", "Zone I", "", "Présidente", "CP", "", "", "05/11/2026", 2, "", "", "", ""])
b2 = wb2.create_sheet("Bureaux"); b2.append(U.ENT_BUREAUX)
wb2.save(buf); buf.seek(0)
t = txt(ok(c.post("/uascep/import", data={"fichier": (buf, "uascep.xlsx")}, content_type="multipart/form-data", follow_redirects=True)))
assert "ligne(s) refusée(s)" in t and "hors de votre périmètre" in t
t = txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
assert "BIO Salimatou" in t and "01/11/2028" in t
if JUR:
    t = txt(ok(c.get("/uascep?onglet=UP&entite=UP 3")))
    assert "KPADE Jean" in t and "ADJAÏ Rachidi" in t
    t = txt(ok(c.get("/uascep?onglet=ZONE")))
    assert "SOSSA Ida" in t
    ok(c.post("/uascep/membre/0?niveau=CS&entite=Pobè", data={"ordre": "1", "nom": "HOUNKPE", "prenoms": "Luc", "fonction": "Président"}), 302)
    assert "HOUNKPE Luc" in txt(ok(c.get("/uascep?onglet=CS&entite=Pobè")))
    if ED == "ddemp":
        ok(c.post("/uascep/membre/0?niveau=DDEMP&entite=Plateau", data={"ordre": "1", "nom": "AGBO", "prenoms": "Rose", "fonction": "Présidente"}), 302)
        assert "AGBO Rose" in txt(ok(c.get("/uascep?onglet=DDEMP&entite=Plateau")))
    else:
        ok(c.get("/uascep?onglet=DDEMP"), 404)
else:
    ok(c.get("/uascep?onglet=CS"), 404)
    ok(c.post("/uascep/membre/0?niveau=UP&entite=X", data={"nom": "Z"}), 403)      # UP : lecture seule à l'école
    # réinitialiser tous les bureaux de l'onglet
    ok(c.post("/uascep/reinitialiser", data={"niveau": "ECOLE"}), 302)
    assert "BIO Salimatou" not in txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
    ok(c.post(f"/uascep/membre/0?niveau=ECOLE&entite={ent}", data={"ordre": "1", "nom": "ZINSOU", "prenoms": "Paul", "fonction": "Président"}), 302)
    ok(c.get("/logout"), 302)
    c = app.test_client()
    ok(c.post("/login", data={"identifiant": "ens1", "mdp": "ens-mdp-123"}), 302)
    t = txt(ok(c.get(f"/uascep?onglet=ECOLE&entite={ent}")))
    assert "ZINSOU Paul" in t and "Ajouter un membre" not in t
    ok(c.post(f"/uascep/membre/0?niveau=ECOLE&entite={ent}", data={"nom": "Z"}), 403)
print("OK-" + config.EDITION)
'''


def lancer(edition):
    env = dict(os.environ)
    for k in ("SCOLARIS_EDITION", "SCOLARIS_DATA", "GESCOL_EDITION", "GESCOL_DATA"):
        env.pop(k, None)
    env["EDITION_SCRIPT"] = edition
    env["PREFIXE_SCRIPT"] = PREFIXE
    return subprocess.run([sys.executable, "-c", SCRIPT], cwd=RACINE, env=env, capture_output=True, text=True, timeout=600)


class TestUASCEP(unittest.TestCase):
    def verifier(self, ed):
        r = lancer(ed)
        self.assertIn("OK-" + ed, r.stdout, r.stdout[-3000:] + r.stderr[-6000:])

    def test_privee(self):
        self.verifier("privee")

    def test_publique(self):
        self.verifier("publique")

    def test_cs(self):
        self.verifier("cs")

    def test_ddemp(self):
        self.verifier("ddemp")


if __name__ == "__main__":
    unittest.main()
